"""
ppt_template.py — Fill the user's OWN .pptx/.potx without changing its design
==============================================================================
The template's look is sacred: backgrounds, fonts, colours, positions, logos and
decorations are kept exactly. We only swap text (keeping each paragraph's/run's
original formatting) and pictures (cropped to the original frame).

Two strategies, picked automatically:
  • clone  — the template has sample slides: for every deck slide we clone the
             best-matching sample slide (title / section / content / picture /
             multi-column / closing) and replace its text + pictures.
  • layout — the template only has layouts (typical .potx): new slides are made
             from the matching layouts and their placeholders are filled.
Original sample slides are removed at the end.

Public API:  TemplateFiller(template_path, deck, out_path).build() -> generator
             describe_template(path) -> short string (for chat)
"""
from __future__ import annotations

import copy
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Optional

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.services.ppt_designer import (prepare_image, _best_window, _saliency_profile,
                                       _items, line_h, count_lines)

EMU_IN = 914400
_TPL_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PRS_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
_SKIP_PH = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}
_TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}


def _inch(v) -> float:
    return (v or 0) / EMU_IN


def load_template(path: str) -> Presentation:
    p = Path(path)
    if p.suffix.lower() in (".potx", ".potm"):
        tmp = Path(tempfile.mkdtemp()) / (p.stem + ".pptx")
        with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(_TPL_CT.encode(), _PRS_CT.encode()) \
                               .replace(_TPL_CT.replace("main+xml", "macroEnabled.main+xml").encode(), _PRS_CT.encode())
                zout.writestr(item, data)
        return Presentation(str(tmp))
    return Presentation(str(p))


# ══════════════════════════════════════════════════════════════════════════════
#  ANALYSIS
# ══════════════════════════════════════════════════════════════════════════════
def _walk(shapes, ox=0, oy=0, sx=1.0, sy=1.0):
    """Yield (shape, abs_left, abs_top, abs_w, abs_h) in inches, recursing into groups."""
    for shp in shapes:
        try:
            l, t, w, h = shp.left or 0, shp.top or 0, shp.width or 0, shp.height or 0
        except Exception:
            l = t = w = h = 0
        ax, ay = ox + l * sx, oy + t * sy
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                xfrm = shp._element.grpSpPr.find(qn("a:xfrm"))
                ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
                cx, cy = int(ch_off.get("x")), int(ch_off.get("y"))
                cw, chh = int(ch_ext.get("cx")) or 1, int(ch_ext.get("cy")) or 1
                nsx, nsy = sx * w / cw, sy * h / chh
                yield from _walk(shp.shapes, ax - cx * nsx, ay - cy * nsy, nsx, nsy)
            except Exception:
                yield from _walk(shp.shapes, ax, ay, sx, sy)
            continue
        yield shp, ax / EMU_IN, ay / EMU_IN, w * sx / EMU_IN, h * sy / EMU_IN


def _ph_type(shp):
    try:
        return shp.placeholder_format.type if shp.is_placeholder else None
    except Exception:
        return None


def _max_size(shp) -> float:
    best = 0.0
    try:
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    best = max(best, r.font.size.pt)
    except Exception:
        pass
    return best


def _analyze_slide(slide, idx: int, sw: float, sh: float) -> dict:
    texts, pics, heavy = [], [], False
    for shp, x, y, w, h in _walk(slide.shapes):
        pt = _ph_type(shp)
        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE or pt in (PP_PLACEHOLDER.PICTURE,) or \
                (shp.is_placeholder and shp.__class__.__name__ == "PlaceholderPicture"):
            if w * h >= sw * sh * 0.04:          # ignore logos / icons
                pics.append({"shape": shp, "x": x, "y": y, "w": w, "h": h})
            continue
        if getattr(shp, "has_chart", False) and shp.has_chart:
            heavy = True
            continue
        if shp.shape_type == MSO_SHAPE_TYPE.TABLE or "dgm" in shp._element.xml[:3000]:
            heavy = True
            continue
        if not getattr(shp, "has_text_frame", False) or not shp.has_text_frame or pt in _SKIP_PH:
            continue
        txt = shp.text_frame.text.strip()
        if not txt and not shp.is_placeholder:
            continue
        role = "title" if pt in _TITLE_PH else "subtitle" if pt == PP_PLACEHOLDER.SUBTITLE else "body"
        texts.append({"shape": shp, "x": x, "y": y, "w": w, "h": h, "role": role, "text": txt,
                      "words": len(txt.split()), "size": _max_size(shp), "ph": pt})
    # no title placeholder → biggest font in the top 40% is the title
    if not any(t["role"] == "title" for t in texts):
        cands = [t for t in texts if t["y"] < sh * 0.4 and t["words"] <= 14]
        if cands:
            best = max(cands, key=lambda t: (t["size"] or 0, -t["y"], t["w"]))
            best["role"] = "title"
    bodies = [t for t in texts if t["role"] == "body" and not (len(t["text"]) <= 3 and not t["shape"].is_placeholder)]
    all_text = " ".join(t["text"].lower() for t in texts)
    info = {"idx": idx, "slide": slide, "texts": texts, "bodies": _reading_order(bodies), "pics": pics,
            "heavy": heavy, "n_body": len(bodies),
            "is_title": idx == 0 or any(t["ph"] == PP_PLACEHOLDER.CENTER_TITLE for t in texts),
            "is_closing": bool(re.search(r"thank|questions|q\s*&\s*a|contact us|the end", all_text)),
            "is_section": len(bodies) <= 1 and not pics and any(t["role"] == "title" for t in texts)
            and sum(t["words"] for t in bodies) <= 12}
    info["pairs"] = _pairs(info["bodies"])
    return info


def _reading_order(items):
    items = sorted(items, key=lambda t: (t["y"], t["x"]))
    rows, out = [], []
    for it in items:
        if rows and abs(rows[-1][0]["y"] - it["y"]) < 0.35:
            rows[-1].append(it)
        else:
            rows.append([it])
    for r in rows:
        out += sorted(r, key=lambda t: t["x"])
    return out


def _pairs(bodies):
    """(head_slot, text_slot) pairs: a short text sitting just above another with the same left edge."""
    pairs, used = [], set()
    for a in bodies:
        if id(a) in used or a["words"] > 7:
            continue
        below = [b for b in bodies if b is not a and id(b) not in used and abs(b["x"] - a["x"]) < 0.4
                 and 0 < b["y"] - a["y"] < max(1.6, a["h"] + 0.9) and b["words"] >= a["words"]]
        if below:
            b = min(below, key=lambda b: b["y"])
            pairs.append((a, b))
            used |= {id(a), id(b)}
    return pairs if len(pairs) >= 2 else []


# ══════════════════════════════════════════════════════════════════════════════
#  FORMAT TEMPLATES  (fixed-section formats: SIH idea PPT, college report formats…)
# ══════════════════════════════════════════════════════════════════════════════
_GUIDE = re.compile(r"\b(describe|explain|e\.g\.|details?|mention|analysis|potential|include|list|your|briefly|write|"
                    r"enter|insert|provide|highlight|add here|goes here|lorem|placeholder|type here)\b|\(.{6,}\)", re.I)
_PLACEHOLDER_TITLE = re.compile(r"\b(title|topic|your|heading here|lorem|click to|insert)\b", re.I)
_SKIP_TITLE = re.compile(r"\b(instructions?|guidelines?|important notes?|read ?me|how to use|note to)\b", re.I)
_FIELDISH = re.compile(r"^\s*([^.!?]{2,50}?)\s*(?:[-–—:]\s*(.*)|\((.*)\)\s*)$")
_SLOT = re.compile(r"\b(team|name|logo|date|college|institute|presenter|member|roll|id)\b", re.I)


def _label(prompt: str) -> str:
    """'Proposed Solution (Describe your idea…)' → 'Proposed Solution'."""
    lab = re.split(r"\s*\(", prompt, 1)[0].strip(" :-–—•❖▪➢►")
    return lab if 1 <= len(lab.split()) <= 9 else ""


def _field_label(line: str) -> str:
    m = _FIELDISH.match(line)
    return m.group(1).strip(" •❖") if m else ""


def classify_box(t: dict, sw: float, sh: float) -> str:
    """title | slot | title_slot | form | prompt | other"""
    if t["role"] == "title":
        return "title"
    txt = t["text"]
    lines = [l.strip() for l in txt.split("\n") if l.strip()]
    small = t["w"] * t["h"] < sw * sh * 0.06
    if t["words"] <= 4 and small and _SLOT.search(txt) and not re.search(r"\btitle\b", txt, re.I):
        return "slot"
    if t["words"] <= 4 and _PLACEHOLDER_TITLE.search(txt):
        return "title_slot"
    if len(lines) >= 2 and sum(1 for l in lines if _FIELDISH.match(l) and len(_field_label(l).split()) <= 6) >= 0.6 * len(lines) \
            and not any(_GUIDE.search(l) and len(l.split()) > 8 for l in lines):
        return "form"
    if lines:
        return "prompt"
    return "other"


def analyze_format(path: str) -> Optional[dict]:
    """If the template is a fixed format (sections with guidance prompts), describe its sections."""
    try:
        prs = load_template(path)
    except Exception:
        return None
    sw, sh = _inch(prs.slide_width), _inch(prs.slide_height)
    sections, guided, forms = [], 0, 0
    for i, s in enumerate(prs.slides):
        info = _analyze_slide(s, i, sw, sh)
        title = next((t["text"] for t in info["texts"] if t["role"] == "title"), "")
        prompts, fields, slots, title_slots = [], [], [], []
        for t in info["texts"]:
            c = classify_box(t, sw, sh)
            if c == "slot":
                slots.append(t["text"].strip())
            elif c == "title_slot":
                title_slots.append(t["text"].strip())
            elif c == "form":
                fields += [_field_label(l) for l in t["text"].split("\n") if l.strip() and _field_label(l)]
            elif c == "prompt":
                prompts += [l.strip() for l in t["text"].split("\n") if l.strip()]
        guided += any(_GUIDE.search(p) for p in prompts)
        forms += bool(fields)
        sections.append({"idx": i, "title": title, "skip": bool(_SKIP_TITLE.search(title)), "prompts": prompts,
                         "fields": fields, "slots": slots, "title_slots": title_slots,
                         "title_is_placeholder": bool(_PLACEHOLDER_TITLE.search(title)) or not title})
    is_format = len(sections) >= 3 and (guided >= 2 or (forms >= 1 and guided >= 1))
    return {"is_format": True, "sections": sections, "n": len(sections)} if is_format else None


def _fit_grow(shp, sh: float, default_size: float):
    """Grow a body box downwards (never past the next shape, e.g. a footer bar), then shrink text if needed."""
    try:
        room = _limit_bottom(shp, sh) - _inch(shp.top)
        size, fam = _run_size_and_font(shp, default_size)
        w = (_inch(shp.width) - 0.2) * 0.96
        need = _text_height(shp, fam, size, w) + 0.15
        if need > _inch(shp.height) and room > _inch(shp.height):
            l, t, wd = shp.left, shp.top, shp.width
            shp.left, shp.top, shp.width = l, t, wd
            shp.height = Emu(int(min(room, need) * EMU_IN))
    except Exception:
        pass
    _shrink_to_fit(shp, default_size)


def fill_form_slide(slide, info: dict, sp: dict, sh: float):
    """Fill one cloned format slide: keep headings, answer each prompt, fill fields/slots."""
    sw = sh * 16 / 9
    sections = sp.get("sections") or []
    fields = sp.get("form_fields") or {}
    slots = sp.get("slots") or {}
    boxes = [(t, classify_box(t, sw, sh)) for t in info["texts"]]
    prompt_boxes = [t for t, c in boxes if c == "prompt"]
    biggest = max(prompt_boxes, key=lambda x: x["w"] * x["h"]) if prompt_boxes else None
    free = [s for s in sections if not (s.get("prompt") or "").strip()]      # user content without prompts
    for t, c in boxes:
        txt = t["text"]
        if c == "title":
            if sp.get("replace_title") and sp.get("title"):
                write_text(t["shape"], [{"runs": [(sp["title"], None)], "role": "head"}], fit=False)
                _shrink_to_fit(t["shape"], 32, sh)
            continue
        if c == "slot":
            val = slots.get(txt.strip())
            if val:
                write_text(t["shape"], [{"runs": [(val, None)], "role": "text"}], fit=False)
                _shrink_to_fit(t["shape"], 14, sh)
            continue
        if c == "title_slot":
            val = sp.get("idea_title") or ""
            if val:
                write_text(t["shape"], [{"runs": [(val, None)], "role": "text"}], fit=False)
                _shrink_to_fit(t["shape"], 24, sh)
            continue
        if c == "form":
            paras = []
            for line in [l for l in txt.split("\n") if l.strip()]:
                lab = _field_label(line)
                val = fields.get(lab, "")
                if lab and val:
                    sep = "–" if re.search(r"[-–—]", line) else ":"
                    paras.append({"runs": [(f"{lab} {sep} ", True), (val, None)], "role": "text"})
                else:
                    paras.append({"runs": [(line.strip(), None)], "role": "text"})    # unknown → left as in the format
            write_text(t["shape"], paras, fit=False)
            _fit_grow(t["shape"], sh, 18)
            continue
        if c != "prompt":
            continue
        lines = [l.strip() for l in txt.split("\n") if l.strip()]
        mine = [s for s in sections if (s.get("prompt") or "").strip() in lines]
        if t is biggest:
            mine = mine + free
        if not mine:
            if _GUIDE.search(txt):
                _clear(t["shape"])                      # guidance nobody answered → remove, keep the box
            continue
        paras = []
        for sec in mine:
            lab = sec.get("label") or _label(sec.get("prompt", ""))
            if lab:
                paras.append({"runs": [(lab, True)], "role": "head"})
            for pt in sec.get("points") or []:
                if isinstance(pt, dict):
                    pt = ((pt.get("head") + ": ") if pt.get("head") else "") + (pt.get("text") or "")
                paras.append({"runs": [(str(pt), None)], "role": "text"})
        write_text(t["shape"], paras, fit=False)
        _fit_grow(t["shape"], sh, 18)
    if sp.get("notes"):
        try:
            slide.notes_slide.notes_text_frame.text = sp["notes"]
        except Exception:
            pass


def describe_template(path: str) -> str:
    try:
        prs = load_template(path)
        return f"{len(prs.slides)} sample slides, {len(prs.slide_layouts)} layouts"
    except Exception as e:
        return f"unreadable ({e})"


# ══════════════════════════════════════════════════════════════════════════════
#  TEXT WRITING (format-preserving)
# ══════════════════════════════════════════════════════════════════════════════
def _protos(txBody):
    """(p_proto, r_proto) for heads and for texts, taken from the template text."""
    ps = txBody.findall(qn("a:p"))
    found = []
    for p in ps:
        r = p.find(qn("a:r"))
        if r is not None:
            found.append((p, r))
    if not found:
        p0 = ps[0] if ps else None
        return (p0, None), (p0, None)
    if len(found) >= 2:
        return found[0], found[1]
    return found[0], found[0]


def _new_p(proto_p):
    if proto_p is None:
        return etree.Element(qn("a:p"))
    p = copy.deepcopy(proto_p)
    for ch in list(p):
        if ch.tag != qn("a:pPr"):
            p.remove(ch)
    return p


def _new_r(proto_r, text: str, bold: Optional[bool]):
    if proto_r is not None:
        r = copy.deepcopy(proto_r)
        rpr = r.find(qn("a:rPr"))
    else:
        r = etree.Element(qn("a:r"))
        rpr = etree.SubElement(r, qn("a:rPr"))
        rpr.set("lang", "en-US")
        etree.SubElement(r, qn("a:t"))
    if rpr is None:
        rpr = etree.Element(qn("a:rPr"))
        r.insert(0, rpr)
    if bold is True:
        rpr.set("b", "1")
    elif bold is False and proto_r is None:
        pass
    t = r.find(qn("a:t"))
    t.text = text
    return r


def write_text(shp, paras: list, fit: bool = True, default_size: float = 18):
    """paras: list of {"runs": [(text, bold|None)], "role": "head"|"text"}."""
    tf = shp.text_frame
    txBody = tf._txBody
    (hp, hr), (tp, tr) = _protos(txBody)
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    if not paras:
        paras = [{"runs": [("", None)], "role": "text"}]
    for para in paras:
        pp, rr = (hp, hr) if para.get("role") == "head" else (tp, tr)
        p = _new_p(pp)
        for text, bold in para["runs"]:
            p.append(_new_r(rr, text, bold))
        txBody.append(p)
    if fit:
        _shrink_to_fit(shp, default_size)


def _lst_attr(el, attr: str, tag: str = "a:defRPr"):
    """First `attr` found on lvl1 defRPr of an lstStyle / txStyles element."""
    if el is None:
        return None
    for lvl in el:
        if not lvl.tag.endswith("lvl1pPr"):
            continue
        d = lvl.find(qn(tag))
        if d is not None:
            if attr == "sz" and d.get("sz"):
                return int(d.get("sz")) / 100
            if attr == "latin":
                lat = d.find(qn("a:latin"))
                if lat is not None and lat.get("typeface"):
                    return lat.get("typeface")
    return None


def _theme_font(shp, minor: bool = True) -> Optional[str]:
    try:
        master = shp.part.slide_layout.slide_master if hasattr(shp.part, "slide_layout") else None
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                root = etree.fromstring(rel.target_part.blob)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                f = root.find(f".//a:fontScheme/a:{'minorFont' if minor else 'majorFont'}/a:latin", ns)
                return f.get("typeface") if f is not None else None
    except Exception:
        return None


def _run_size_and_font(shp, default_size):
    """Effective font size/family: runs → shape lstStyle → layout/master placeholder → master txStyles."""
    size, fam = None, None
    for r in shp._element.iter(qn("a:rPr")):
        if size is None and r.get("sz"):
            size = int(r.get("sz")) / 100
        lat = r.find(qn("a:latin"))
        if fam is None and lat is not None and lat.get("typeface") and not lat.get("typeface", "").startswith("+"):
            fam = lat.get("typeface")
    try:
        chain = [shp]
        node = shp
        while getattr(node, "is_placeholder", False) and hasattr(node, "_base_placeholder") and len(chain) < 3:
            node = node._base_placeholder
            if node is None:
                break
            chain.append(node)
        for n in chain:
            lst = n._element.find(".//" + qn("a:lstStyle"))
            size = size or _lst_attr(lst, "sz")
            if fam is None:
                f = _lst_attr(lst, "latin")
                fam = f if f and not f.startswith("+") else None
            if size is None:
                for r in n._element.iter(qn("a:endParaRPr")):
                    if r.get("sz"):
                        size = int(r.get("sz")) / 100
                        break
        pt = _ph_type(shp)
        master = shp.part.slide_layout.slide_master
        tx = master._element.find(qn("p:txStyles"))
        if tx is not None and size is None:
            key = "p:titleStyle" if pt in _TITLE_PH else "p:bodyStyle" if shp.is_placeholder else "p:otherStyle"
            size = _lst_attr(tx.find(qn(key)), "sz")
        if fam is None:
            fam = _theme_font(shp, minor=pt not in _TITLE_PH)
    except Exception:
        pass
    return size or default_size, fam or "Calibri"


def _limit_bottom(shp, slide_h: float) -> float:
    """Lowest y the box may reach: the first other shape below it that overlaps horizontally (footer bars…)."""
    try:
        top, left, right = _inch(shp.top), _inch(shp.left), _inch(shp.left) + _inch(shp.width)
        lim = slide_h - 0.25
        for o in shp.part.slide.shapes:
            if o.shape_id == shp.shape_id or o.top is None or o.width is None:
                continue
            ot, ol, orr = _inch(o.top), _inch(o.left), _inch(o.left) + _inch(o.width)
            if ot > top + _inch(shp.height) * 0.6 and ol < right and orr > left:
                lim = min(lim, ot - 0.05)
        return lim
    except Exception:
        return slide_h - 0.4


def _spacing(p, size: float):
    """(line-spacing multiple | ('pts', v), space before (in), space after (in)) of a paragraph."""
    ls, bef, aft = 1.0, 0.0, 0.0
    pPr = p._p.find(qn("a:pPr"))
    if pPr is None:
        return ls, bef, aft
    ln = pPr.find(qn("a:lnSpc"))
    if ln is not None:
        pct, pts = ln.find(qn("a:spcPct")), ln.find(qn("a:spcPts"))
        if pct is not None:
            ls = int(pct.get("val")) / 100000
        elif pts is not None:
            ls = ("pts", int(pts.get("val")) / 100)
    vals = []
    for tag in ("a:spcBef", "a:spcAft"):
        el, v = pPr.find(qn(tag)), 0.0
        if el is not None:
            sp, pc_ = el.find(qn("a:spcPts")), el.find(qn("a:spcPct"))
            if sp is not None:
                v = int(sp.get("val")) / 100 / 72
            elif pc_ is not None:
                v = int(pc_.get("val")) / 100000 * size * 1.2 / 72
        vals.append(v)
    return ls, vals[0], vals[1]


def _text_height(shp, fam: str, base: float, w: float, scale: float = 1.0) -> float:
    total = 0.0
    for p in shp.text_frame.paragraphs:
        runs = [(r.text, bool(r.font.bold)) for r in p.runs] or [("", False)]
        size = max([r.font.size.pt if r.font.size else base for r in p.runs] or [base]) * scale
        ls, bef, aft = _spacing(p, size)
        n = count_lines(runs, fam, size, w)
        total += (n * ls[1] * scale / 72 if isinstance(ls, tuple) else n * line_h(fam, size, ls)) + bef + aft
    return total


def _shrink_to_fit(shp, default_size, slide_h: float = None):
    """Scale all runs proportionally (keeps heading/body hierarchy) until the text fits its box."""
    try:
        w = (_inch(shp.width) - 0.2) * 0.96
        top = _inch(shp.top)
        h = _inch(shp.height) - 0.1
        bp = shp.text_frame._txBody.find(qn("a:bodyPr"))
        if bp is not None and bp.find(qn("a:spAutoFit")) is not None and slide_h:
            h = max(h, _limit_bottom(shp, slide_h) - top - 0.1)       # auto-growing box: stop at the next shape
        if w <= 0.3 or h <= 0.2:
            return
        base, fam = _run_size_and_font(shp, default_size)
        if _text_height(shp, fam, base, w) <= h:
            return
        smallest = min([r.font.size.pt if r.font.size else base for p in shp.text_frame.paragraphs for r in p.runs] or [base])
        scale = 1.0
        while scale > 0.3 and smallest * scale > 9 and _text_height(shp, fam, base, w, scale) > h:
            scale -= 0.04
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(max(8, round((r.font.size.pt if r.font.size else base) * scale * 2) / 2))
    except Exception as e:
        print(f"[ppt_template] fit skipped: {e}")


def _clear(shp):
    try:
        write_text(shp, [], fit=False)
    except Exception:
        pass


# ══════════════════════════════════════════════════════════════════════════════
#  PICTURES
# ══════════════════════════════════════════════════════════════════════════════
def _replace_picture(slide, pic, img_path: str):
    path, aspect = prepare_image(img_path)
    if not path:
        return
    if pic.__class__.__name__ == "SlidePlaceholder" or (pic.is_placeholder and not hasattr(pic, "image")):
        try:
            pic.insert_picture(path)
            return
        except Exception:
            pass
    _, rId = slide.part.get_or_add_image_part(path)
    blip = pic._element.find(".//" + qn("a:blip"))
    if blip is None:
        return
    blip.set(qn("r:embed"), rId)
    fw, fh = pic.width or 1, pic.height or 1
    box_a = fw / fh
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0.0
    if aspect > box_a:
        keep = box_a / aspect
        st = _best_window(_saliency_profile(path, "x"), keep)
        pic.crop_left, pic.crop_right = st, max(0.0, 1 - keep - st)
    elif aspect < box_a:
        keep = aspect / box_a
        st = min(_best_window(_saliency_profile(path, "y"), keep), (1 - keep) * 0.45)
        pic.crop_top, pic.crop_bottom = st, max(0.0, 1 - keep - st)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE CLONING
# ══════════════════════════════════════════════════════════════════════════════
_R_ATTRS = [qn("r:embed"), qn("r:link"), qn("r:id"), qn("r:pict")]


def clone_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rmap = {}
    for rId, rel in list(src.part.rels.items()):
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        try:
            if rel.is_external:
                rmap[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rmap[rId] = new.part.relate_to(rel.target_part, rel.reltype)
        except Exception:
            pass
    src_cSld = src._element.find(qn("p:cSld"))
    new_cSld = new._element.find(qn("p:cSld"))
    bg = src_cSld.find(qn("p:bg"))
    if bg is not None:
        new_cSld.insert(0, copy.deepcopy(bg))
    spTree = new.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        spTree.append(copy.deepcopy(el))
    for el in [new._element]:
        for node in el.iter():
            for a in _R_ATTRS:
                v = node.get(a)
                if v and v in rmap:
                    node.set(a, rmap[v])
    # keep transition / timing-free: copy clrMapOvr if present
    cmo = src._element.find(qn("p:clrMapOvr"))
    if cmo is not None and new._element.find(qn("p:clrMapOvr")) is None:
        new._element.append(copy.deepcopy(cmo))
    return new


def delete_slide(prs, slide):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        if prs.part.related_part(sldId.rId) is slide.part:
            rId = sldId.rId
            sldIdLst.remove(sldId)
            try:
                prs.part.drop_rel(rId)
            except Exception:
                prs.part.rels.pop(rId)
            return


# ══════════════════════════════════════════════════════════════════════════════
#  CONTENT → UNITS
# ══════════════════════════════════════════════════════════════════════════════
def _units(sp: dict) -> tuple[list[dict], str]:
    """Deck slide → ([{head, text}], lead)."""
    kind = sp.get("kind")
    lead = sp.get("body") or ""
    if kind == "stats":
        return [{"head": s["value"], "text": " — ".join(x for x in (s.get("label"), s.get("desc")) if x)}
                for s in sp.get("stats") or []], lead
    if kind == "timeline":
        return [{"head": " — ".join(x for x in (st.get("date"), st.get("head")) if x), "text": st.get("text", "")}
                for st in sp.get("steps") or []], lead
    if kind == "comparison":
        return [{"head": c.get("heading", ""), "text": "\n".join(c.get("points") or [])}
                for c in sp.get("columns") or []], lead
    if kind == "quote":
        q = sp.get("quote") or {}
        return [{"head": "", "text": f"“{q.get('text', '')}”" + (f"\n— {q['author']}" if q.get("author") else "")}], ""
    if kind == "table" and sp.get("table"):
        tb = sp["table"]
        rows = [tb.get("header") or []] + (tb.get("rows") or [])
        return [{"head": "", "text": " | ".join(map(str, r))} for r in rows], lead
    if kind == "chart" and sp.get("chart"):
        ch = sp["chart"]
        vals = (ch.get("series") or [{}])[0].get("values") or ch.get("values") or []
        its = [{"head": str(l), "text": f"{v:g}{ch.get('unit', '')}" if isinstance(v, (int, float)) else str(v)}
               for l, v in zip(ch.get("labels") or [], vals)]
        return its + [{"head": "", "text": i["text"]} for i in _items(sp)], lead
    return [{"head": i.get("head", ""), "text": i.get("text", "")} for i in _items(sp)], lead


def _paras_for(units, lead="", stacked=False):
    paras = []
    if lead:
        paras.append({"runs": [(lead, None)], "role": "text"})
    for u in units:
        if stacked and u["head"] and u["text"]:
            paras.append({"runs": [(u["head"], True)], "role": "head"})
            for line in u["text"].split("\n"):
                paras.append({"runs": [(line, None)], "role": "text"})
        elif u["head"] and u["text"]:
            first, *rest = u["text"].split("\n")
            paras.append({"runs": [(u["head"] + ": ", True), (first, None)], "role": "text"})
            paras += [{"runs": [(r, None)], "role": "text"} for r in rest]
        else:
            for line in (u["head"] or u["text"]).split("\n"):
                paras.append({"runs": [(line, None)], "role": "text"})
    return paras


# ══════════════════════════════════════════════════════════════════════════════
#  FILLER
# ══════════════════════════════════════════════════════════════════════════════
class TemplateFiller:
    def __init__(self, template_path: str, deck: dict, out_path: str):
        self.path = template_path
        self.deck = deck
        self.out = out_path

    def build(self):
        prs = load_template(self.path)
        sw, sh = _inch(prs.slide_width), _inch(prs.slide_height)
        originals = list(prs.slides)
        infos = [_analyze_slide(s, i, sw, sh) for i, s in enumerate(originals)]
        slides = self.deck.get("slides", [])
        mode = "format" if self.deck.get("format_mode") else "clone" if infos else "layout"
        yield f"🧩 Template: {len(originals)} sample slide(s), {len(prs.slide_layouts)} layout(s) → {mode} mode"
        prev = None
        for i, sp in enumerate(slides):
            try:
                tpl = sp.get("tpl")
                if isinstance(tpl, int) and 0 <= tpl < len(infos):
                    new = clone_slide(prs, infos[tpl]["slide"])
                    fill_form_slide(new, _analyze_slide(new, tpl, sw, sh), sp, sh)
                elif mode in ("clone", "format"):
                    info = self._choose(infos, sp, i, len(slides), prev)
                    prev = info["idx"]
                    new = clone_slide(prs, info["slide"])
                    self._fill_clone(prs, new, _analyze_slide(new, info["idx"], sw, sh), sp)
                else:
                    self._fill_layout(prs, sp, i)
                yield f"  ✓ Slide {i + 1}/{len(slides)}: {sp.get('title', '')[:60]}"
            except Exception as e:
                yield f"  ⚠️ Slide {i + 1} fell back to a plain layout ({e})"
                try:
                    self._fill_layout(prs, sp, i)
                except Exception:
                    pass
        for s in originals:
            delete_slide(prs, s)
        prs.save(self.out)

    # ── choose a sample slide ──────────────────────────────────────────────
    def _choose(self, infos, sp, i, n, prev):
        kind = sp.get("kind", "content")
        has_img = bool(sp.get("images"))
        units, lead = _units(sp)
        m = len(units)
        light = [x for x in infos if not x["heavy"]] or infos

        if kind == "title":
            c = [x for x in light if x["is_title"]]
            return c[0] if c else light[0]
        if kind == "closing":
            c = [x for x in light if x["is_closing"]] or [x for x in light if x["idx"] == len(infos) - 1 and x["n_body"] <= 1] \
                or [x for x in light if x["is_title"]]
            return c[-1] if c else light[-1]
        if kind == "section":
            c = [x for x in light if x["is_section"] and not x["is_title"]]
            if c:
                return c[0]

        def score(x):
            s = 0.0
            if x["is_title"] and len(light) > 1:
                s -= 6
            if x["is_closing"] and len(light) > 2:
                s -= 6
            if has_img:
                s += 3 if x["pics"] else -1
            elif x["pics"]:
                s -= 1.5
            if x["pairs"]:
                s += 3 if len(x["pairs"]) == m else (1.5 if kind in ("cards", "process", "timeline", "stats", "comparison") else -1)
            elif x["n_body"] >= 2:
                if kind == "comparison" and x["n_body"] == len(sp.get("columns") or []):
                    s += 3
                elif m == x["n_body"] or (m >= 2 * x["n_body"] and kind in ("cards", "content", "process")):
                    s += 1
                else:
                    s -= 1
            elif x["n_body"] == 1:
                s += 2 if kind in ("content", "agenda", "quote", "table", "chart") or m > 4 else 1
            else:
                s -= 3
            if x["idx"] == prev:
                s -= 0.4            # variety matters less than fit in someone else's template
            return s

        return max(light, key=score)

    # ── fill a cloned slide ───────────────────────────────────────────────
    def _fill_clone(self, prs, slide, info, sp):
        kind = sp.get("kind", "content")
        units, lead = _units(sp)
        filled = set()
        for t in info["texts"]:
            if t["role"] == "title":
                write_text(t["shape"], [{"runs": [(sp.get("title", ""), None)], "role": "head"}], default_size=36)
                filled.add(id(t["shape"]))
                break
        subs = [t for t in info["texts"] if t["role"] == "subtitle"]
        bodies = [b for b in info["bodies"] if b["role"] == "body"]
        if kind in ("title", "closing", "section"):
            sub = sp.get("subtitle") or ""
            meta = sp.get("body") or " · ".join(((u["head"] + ": ") if u["head"] and u["text"] else "") +
                                                 (u["text"] or u["head"]) for u in units)
            slots = subs + bodies
            if slots and sub:
                write_text(slots[0]["shape"], [{"runs": [(sub, None)], "role": "text"}], default_size=20)
                filled.add(id(slots[0]["shape"]))
                slots = slots[1:]
            if slots and meta:
                write_text(slots[0]["shape"], [{"runs": [(meta, None)], "role": "text"}], default_size=16)
                filled.add(id(slots[0]["shape"]))
            elif not slots and (sub or meta) and not subs:
                pass
        else:
            pairs = [(a, b) for a, b in info["pairs"]]
            if pairs and len(units) >= 2:
                for k, (hs, ts) in enumerate(pairs):
                    if k < len(units):
                        u = units[k]
                        extra = units[len(pairs):] if k == len(pairs) - 1 else []
                        write_text(hs["shape"], [{"runs": [(u["head"] or u["text"].split(".")[0][:40], None)], "role": "head"}],
                                   default_size=20)
                        body_paras = _paras_for([{"head": "", "text": u["text"] if u["head"] else ""}] +
                                                [{"head": e["head"], "text": e["text"]} for e in extra])
                        write_text(ts["shape"], [p for p in body_paras if any(r[0] for r in p["runs"])] or
                                   [{"runs": [("", None)], "role": "text"}], default_size=16)
                        filled |= {id(hs["shape"]), id(ts["shape"])}
                rest = [b for b in bodies if id(b["shape"]) not in filled]
                if lead and rest:
                    write_text(rest[0]["shape"], [{"runs": [(lead, None)], "role": "text"}], default_size=16)
                    filled.add(id(rest[0]["shape"]))
            elif len(bodies) >= 2 and (len(units) == len(bodies) or kind == "comparison"):
                for k, b in enumerate(bodies):
                    if k < len(units):
                        write_text(b["shape"], _paras_for([units[k]], stacked=True), default_size=18)
                        filled.add(id(b["shape"]))
            elif len(bodies) >= 2 and len(units) > len(bodies):
                # spread the items across the columns instead of leaving some empty
                per = -(-len(units) // len(bodies))
                for k, b in enumerate(bodies):
                    part = units[k * per:(k + 1) * per]
                    if part or (k == 0 and lead):
                        write_text(b["shape"], _paras_for(part, lead if k == 0 else ""), default_size=18)
                        filled.add(id(b["shape"]))
            elif bodies:
                main = max(bodies, key=lambda b: b["w"] * b["h"])
                others = [b for b in bodies if b is not main]
                if lead and others and len(lead.split()) > 12:
                    write_text(others[0]["shape"], [{"runs": [(lead, None)], "role": "text"}], default_size=16)
                    filled.add(id(others[0]["shape"]))
                    lead_in_main = ""
                else:
                    lead_in_main = lead
                write_text(main["shape"], _paras_for(units, lead_in_main), default_size=18)
                filled.add(id(main["shape"]))
                if kind == "table" and sp.get("table"):
                    self._table_over(slide, main, sp["table"])
            else:
                self._free_text(slide, info, units, lead)
        # pictures
        imgs = list(sp.get("images") or [])
        pics = sorted(info["pics"], key=lambda p: -p["w"] * p["h"])
        for pic, img in zip(pics, imgs):
            _replace_picture(slide, pic["shape"], img)
        if imgs and not pics:
            self._add_picture_beside(slide, info, imgs[0])
        # clear leftover sample text (keep tiny labels like "01", icons, numbers)
        for t in info["texts"]:
            if id(t["shape"]) in filled or t["role"] == "title":
                continue
            if len(t["text"]) > 3 or t["shape"].is_placeholder:
                _clear(t["shape"])
        if sp.get("notes"):
            try:
                slide.notes_slide.notes_text_frame.text = sp["notes"]
            except Exception:
                pass

    def _table_over(self, slide, body, table):
        header = [str(x) for x in table.get("header") or []]
        rows = [[str(c) for c in r] for r in table.get("rows") or []]
        if not header:
            return
        _clear(body["shape"])
        n = len(header)
        shape = slide.shapes.add_table(len(rows) + 1, n, Emu(int(body["x"] * EMU_IN)), Emu(int(body["y"] * EMU_IN)),
                                       Emu(int(body["w"] * EMU_IN)), Emu(int(min(body["h"], 0.4 * (len(rows) + 1)) * EMU_IN)))
        for r, vals in enumerate([header] + rows):
            for c in range(n):
                shape.table.cell(r, c).text = vals[c] if c < len(vals) else ""

    def _free_text(self, slide, info, units, lead):
        title = next((t for t in info["texts"] if t["role"] == "title"), None)
        x = title["x"] if title else 0.8
        y = (title["y"] + title["h"] + 0.3) if title else 1.6
        w = title["w"] if title else 10
        tb = slide.shapes.add_textbox(Emu(int(x * EMU_IN)), Emu(int(y * EMU_IN)), Emu(int(w * EMU_IN)),
                                      Emu(int(max(1.0, 6.8 - y) * EMU_IN)))
        tb.text_frame.word_wrap = True
        write_text(tb, _paras_for(units, lead), default_size=18)

    def _add_picture_beside(self, slide, info, img):
        path, a = prepare_image(img)
        if not path:
            return
        bodies = info["bodies"]
        if bodies:
            main = max(bodies, key=lambda b: b["w"] * b["h"])
            shp = main["shape"]
            if shp.width and main["w"] > 4:
                new_w = main["w"] * 0.56
                # an inherited placeholder has no own xfrm: setting only width would zero its offset
                l, t, h = shp.left, shp.top, shp.height
                shp.left, shp.top, shp.height = l, t, h
                shp.width = Emu(int(new_w * EMU_IN))
                bx, by, bw, bh = main["x"] + new_w + 0.3, main["y"], main["w"] - new_w - 0.3, main["h"]
                _shrink_to_fit(shp, 18)
            else:
                bx, by, bw, bh = main["x"], main["y"] + main["h"] + 0.2, main["w"], 2.0
        else:
            bx, by, bw, bh = 7.0, 1.6, 5.5, 4.8
        if a > bw / bh:
            w, h = bw, bw / a
        else:
            w, h = bh * a, bh
        slide.shapes.add_picture(path, Emu(int((bx + (bw - w) / 2) * EMU_IN)), Emu(int((by + (bh - h) / 2) * EMU_IN)),
                                 Emu(int(w * EMU_IN)), Emu(int(h * EMU_IN)))

    # ── layout mode ──────────────────────────────────────────────────────────
    def _pick_layout(self, prs, kind, has_img, n_cols):
        def phs(l):
            return [p.placeholder_format.type for p in l.placeholders]
        layouts = list(prs.slide_layouts)
        name = lambda l: (l.name or "").lower()
        if kind == "title":
            c = [l for l in layouts if PP_PLACEHOLDER.CENTER_TITLE in phs(l)]
            if c:
                return c[0]
        if kind == "section":
            c = [l for l in layouts if "section" in name(l)]
            if c:
                return c[0]
        if has_img:
            c = [l for l in layouts if PP_PLACEHOLDER.PICTURE in phs(l)]
            if c:
                return c[0]
        bodies = lambda l: sum(1 for t in phs(l) if t in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT))
        if kind == "comparison" and n_cols >= 2:
            c = [l for l in layouts if bodies(l) >= 2]
            if c:
                return c[0]
        c = [l for l in layouts if bodies(l) == 1 and any(t in _TITLE_PH for t in phs(l))]
        if c:
            return c[0]
        c = [l for l in layouts if any(t in _TITLE_PH for t in phs(l))]
        return c[0] if c else layouts[0]

    def _fill_layout(self, prs, sp, i):
        kind = sp.get("kind", "content")
        units, lead = _units(sp)
        imgs = sp.get("images") or []
        layout = self._pick_layout(prs, kind, bool(imgs), len(sp.get("columns") or []))
        slide = prs.slides.add_slide(layout)
        bodies, pics = [], []
        for ph in slide.placeholders:
            t = ph.placeholder_format.type
            if t in _TITLE_PH:
                ph.text_frame.text = sp.get("title", "")
            elif t == PP_PLACEHOLDER.SUBTITLE:
                ph.text_frame.text = sp.get("subtitle") or lead or ""
            elif t == PP_PLACEHOLDER.PICTURE:
                pics.append(ph)
            elif t in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                bodies.append(ph)
        bodies.sort(key=lambda p: (round(_inch(p.top) * 2), p.left or 0))   # rows (½" bands), then left→right
        if kind == "comparison" and len(bodies) >= 2:
            for b, u in zip(bodies, units):
                write_text(b, _paras_for([u], stacked=True), default_size=20)
        elif bodies:
            text_units = units if kind not in ("title", "section") else []
            txt_lead = lead if kind not in ("title",) else (lead if not sp.get("subtitle") else "")
            write_text(bodies[0], _paras_for(text_units, txt_lead), default_size=20)
            for b in bodies[1:]:
                if imgs and not pics:
                    pics.append(b)
                else:
                    b._element.getparent().remove(b._element)
        for ph, img in zip(pics, imgs):
            path, _ = prepare_image(img)
            if path:
                try:
                    ph.insert_picture(path)
                except Exception:
                    pass
        for ph in list(slide.placeholders):              # no "Click to add text" leftovers
            try:
                empty = ph.has_text_frame and not ph.text_frame.text.strip()
                if empty and ph.placeholder_format.type not in _TITLE_PH:
                    ph._element.getparent().remove(ph._element)
            except Exception:
                pass
        if sp.get("notes"):
            slide.notes_slide.notes_text_frame.text = sp["notes"]
        return slide
