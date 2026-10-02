"""
ppt_designer.py — Adaptive Design Engine (PPT v6)
==================================================
Renders a *deck spec* (plain JSON dict, see ppt_content.normalize_deck) into a
.pptx with no fixed templates. Every slide is laid out at render time from:

  • real text measurement (PIL + the actual Windows TTF files) → the largest
    font size that fits, so sparse slides get big type and dense slides stay
    tidy instead of overflowing;
  • the content's shape (item count, words per item, numbers, dates, groups);
  • each image's real aspect ratio → bleed panel / framed / band / justified
    gallery, chosen so images are cropped as little as possible, and when a
    crop is needed a saliency window keeps the interesting part in frame;
  • a 12-col grid with fixed margins, so everything shares the same edges.

Public API:
    THEMES, resolve_theme(name_or_palette, prompt, purpose) -> theme dict
    DeckRenderer(deck, out_path).render() -> generator of progress strings
    prepare_image(path) -> (usable_path, aspect)
"""
from __future__ import annotations

import hashlib
import math
import re
from functools import lru_cache
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt
from lxml import etree

_ROOT = Path(__file__).resolve().parents[2]
_CACHE_DIR = _ROOT / "data" / "ppt_decks" / "img_cache"

EMU_IN = 914400
SW, SH = 13.333, 7.5          # 16:9 slide, inches
MX = 0.75                     # left/right margin
MT = 0.55                     # top margin
MB = 0.62                     # bottom margin (footer lives inside it)
GAP = 0.3                     # standard gutter


def E(x: float) -> int:
    return int(round(x * EMU_IN))


# ══════════════════════════════════════════════════════════════════════════════
#  THEMES
# ══════════════════════════════════════════════════════════════════════════════
# bg / surface / surface2 / text / muted / accent / accent2 / accent3 / line
THEMES: dict[str, dict] = {
    "minimal_light": dict(name="Minimal Light", bg="FFFFFF", surface="F4F6FA", surface2="E9EDF4",
                          text="111827", muted="5B6475", accent="4F46E5", accent2="0EA5E9", accent3="F59E0B",
                          line="E2E6EE", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "editorial":     dict(name="Editorial", bg="FAF7F2", surface="FFFFFF", surface2="F1EBE1",
                          text="1F1B16", muted="6B6258", accent="C2410C", accent2="0F766E", accent3="B45309",
                          line="E6DDD0", head_font="Georgia", body_font="Segoe UI"),
    "slate_corporate": dict(name="Slate Corporate", bg="FFFFFF", surface="F1F5F9", surface2="E2E8F0",
                          text="0F172A", muted="475569", accent="0F4C81", accent2="14B8A6", accent3="F97316",
                          line="E2E8F0", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "ocean_light":   dict(name="Ocean Light", bg="F5FAFF", surface="FFFFFF", surface2="E3F0FC",
                          text="0B2545", muted="4A6180", accent="0284C7", accent2="1D4ED8", accent3="F97316",
                          line="D9E6F2", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "forest":        dict(name="Forest", bg="F5F8F3", surface="FFFFFF", surface2="E6EFE2",
                          text="14281D", muted="52665A", accent="15803D", accent2="CA8A04", accent3="0E7490",
                          line="DCE6D6", head_font="Georgia", body_font="Segoe UI"),
    "medical":       dict(name="Medical", bg="F5FBFB", surface="FFFFFF", surface2="E0F2F1",
                          text="0B3B3C", muted="4B6B6C", accent="0D9488", accent2="2563EB", accent3="F59E0B",
                          line="D3ECEA", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "lavender":      dict(name="Lavender", bg="FBFAFF", surface="FFFFFF", surface2="EEEAFD",
                          text="1E1B4B", muted="5B5A7A", accent="7C3AED", accent2="DB2777", accent3="0EA5E9",
                          line="E4DFF7", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "mono_bold":     dict(name="Mono Bold", bg="F5F5F4", surface="FFFFFF", surface2="E7E5E4",
                          text="0C0A09", muted="57534E", accent="0C0A09", accent2="EAB308", accent3="DC2626",
                          line="E7E5E4", head_font="Bahnschrift", body_font="Segoe UI"),
    "midnight":      dict(name="Midnight", bg="0B1020", surface="141B2E", surface2="1C2540",
                          text="F1F5F9", muted="94A3B8", accent="38BDF8", accent2="A78BFA", accent3="34D399",
                          line="253052", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "neon_pitch":    dict(name="Neon Pitch", bg="09090D", surface="15151C", surface2="1E1E28",
                          text="FAFAFA", muted="A1A1AA", accent="A3E635", accent2="22D3EE", accent3="F472B6",
                          line="2A2A36", head_font="Bahnschrift", body_font="Segoe UI"),
    "cosmic":        dict(name="Cosmic", bg="0B0B1A", surface="15152B", surface2="1E1E3A",
                          text="EEF0FF", muted="9AA0C8", accent="818CF8", accent2="F0ABFC", accent3="67E8F9",
                          line="2A2A4A", head_font="Segoe UI Semibold", body_font="Segoe UI"),
    "ember":         dict(name="Ember", bg="17120F", surface="231B16", surface2="2E241D",
                          text="FFF7ED", muted="C8B6A6", accent="FB923C", accent2="F43F5E", accent3="FACC15",
                          line="3A2E25", head_font="Georgia", body_font="Segoe UI"),
    "emerald_dark":  dict(name="Emerald Dark", bg="06140E", surface="0D2219", surface2="133024",
                          text="ECFDF5", muted="8DB8A5", accent="34D399", accent2="FBBF24", accent3="38BDF8",
                          line="1C3A2D", head_font="Segoe UI Semibold", body_font="Segoe UI"),
}
_DARK_THEMES = ["midnight", "neon_pitch", "cosmic", "ember", "emerald_dark"]
_LIGHT_THEMES = ["minimal_light", "editorial", "slate_corporate", "ocean_light", "forest", "medical", "lavender", "mono_bold"]

_TOPIC_THEMES = [
    (("hackathon", "startup", "pitch", "investor", "mvp", "demo day", "product launch"), "neon_pitch", "minimal_light"),
    (("ai", "machine learning", "deep learning", "neural", "llm", "cyber", "blockchain", "cloud", "software",
      "algorithm", "data science", "robot", "quantum", "tech"), "midnight", "minimal_light"),
    (("space", "astronomy", "galaxy", "universe", "nasa", "rocket", "physics"), "cosmic", "ocean_light"),
    (("medical", "health", "biology", "clinical", "hospital", "patient", "dna", "drug", "vaccine", "disease"), "emerald_dark", "medical"),
    (("finance", "investment", "banking", "revenue", "market", "stock", "economy", "business", "corporate",
      "quarterly", "sales", "strategy", "management"), "midnight", "slate_corporate"),
    (("environment", "climate", "green", "sustainab", "ecology", "nature", "forest", "renewable", "carbon", "agricultur"), "emerald_dark", "forest"),
    (("history", "literature", "culture", "philosophy", "art", "poetry", "heritage", "war"), "ember", "editorial"),
    (("design", "creative", "brand", "marketing", "fashion", "ux", "social media"), "cosmic", "lavender"),
    (("ocean", "water", "marine", "sea", "travel", "tourism"), "midnight", "ocean_light"),
    (("food", "cooking", "restaurant", "coffee", "energy", "solar", "sport", "fitness"), "ember", "editorial"),
    (("education", "school", "university", "learning", "student", "research", "science"), "midnight", "ocean_light"),
]


def _lum(h: str) -> float:
    h = h.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def _mix(a: str, b: str, t: float) -> str:
    a, b = a.lstrip("#"), b.lstrip("#")
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def _contrast(a: str, b: str) -> float:
    la, lb = _lum(a), _lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def on_color(fill: str) -> str:
    """Readable text colour on top of `fill`."""
    return "FFFFFF" if _contrast(fill, "FFFFFF") >= _contrast(fill, "111111") else "111111"


def _finish_theme(t: dict) -> dict:
    t = dict(t)
    t.setdefault("head_font", "Segoe UI Semibold")
    t.setdefault("body_font", "Segoe UI")
    t["dark"] = _lum(t["bg"]) < 0.2
    t.setdefault("line", _mix(t["bg"], t["text"], 0.12))
    t.setdefault("surface2", _mix(t["surface"], t["text"], 0.06))
    # Accent used for text must stay readable on the background.
    acc_txt = t["accent"]
    for _ in range(6):
        if _contrast(acc_txt, t["bg"]) >= 3.0:
            break
        acc_txt = _mix(acc_txt, t["text"], 0.25)
    t["accent_text"] = acc_txt
    t["on_accent"] = on_color(t["accent"])
    t["series"] = [t["accent"], t["accent2"], t["accent3"], _mix(t["accent"], t["text"], 0.45),
                   _mix(t["accent2"], t["bg"], 0.35), _mix(t["accent3"], t["text"], 0.35)]
    return t


def theme_from_palette(p: dict, name: str = "custom") -> dict:
    """Convert a legacy ppt_tool palette (bg/card/text/sub/ac1..) into a theme."""
    g = lambda k, d: (str(p.get(k) or d)).lstrip("#")[:6] or d
    bg = g("bg", "0B1020")
    return _finish_theme(dict(
        name=p.get("name", name), bg=bg, surface=g("card", _mix(bg, "FFFFFF", 0.06)),
        surface2=g("card2", _mix(bg, "FFFFFF", 0.1)), text=g("text", on_color(bg)), muted=g("sub", "888888"),
        accent=g("ac1", "4F46E5"), accent2=g("ac2", "0EA5E9"), accent3=g("ac3", "F59E0B"),
    ))


def list_themes() -> dict:
    return {k: v["name"] for k, v in THEMES.items()}


def resolve_theme(choice=None, prompt: str = "", purpose: str = "general",
                  legacy_palettes: Optional[dict] = None) -> dict:
    """choice may be a THEMES key, a legacy palette key, a palette dict, or None (auto)."""
    if isinstance(choice, dict):
        return theme_from_palette(choice) if "ac1" in choice else _finish_theme({**THEMES["minimal_light"], **choice})
    if isinstance(choice, str) and choice:
        key = choice.strip().lower().replace(" ", "_").replace("-", "_")
        if key in THEMES:
            return _finish_theme({**THEMES[key], "key": key})
        if legacy_palettes and key in legacy_palettes:
            return theme_from_palette(legacy_palettes[key], key)
    low = (prompt or "").lower()
    # explicit theme name in prompt
    for k, v in THEMES.items():
        if k.replace("_", " ") in low or k in low or v["name"].lower() in low:
            return _finish_theme({**THEMES[k], "key": k})
    want_dark = bool(re.search(r"\bdark\b|\bblack\b|\bnight\b", low))
    want_light = bool(re.search(r"\blight\b|\bwhite\b|\bbright\b|\bminimal\b|\bclean\b", low))
    words = set(re.findall(r"[a-z]+", low))
    pick = None
    for kws, dark, light in _TOPIC_THEMES:
        if any((k in words) if " " not in k and len(k) <= 3 else (k in low) for k in kws):
            pick = (dark, light)
            break
    if pick is None:
        pick = ("midnight", "minimal_light")
    if want_dark:
        key = pick[0]
    elif want_light:
        key = pick[1]
    else:
        key = pick[0] if purpose == "hackathon" else pick[1]
    return _finish_theme({**THEMES[key], "key": key})


# ══════════════════════════════════════════════════════════════════════════════
#  TEXT METRICS  (measure with the real fonts so fitting is honest)
# ══════════════════════════════════════════════════════════════════════════════
_FONT_DIR = Path("C:/Windows/Fonts")
_FONT_FILES = {
    "segoe ui": ("segoeui.ttf", "segoeuib.ttf"),
    "segoe ui semibold": ("seguisb.ttf", "segoeuib.ttf"),
    "segoe ui light": ("segoeuil.ttf", "seguisb.ttf"),
    "segoe ui semilight": ("segoeuisl.ttf", "seguisb.ttf"),
    "georgia": ("georgia.ttf", "georgiab.ttf"),
    "calibri": ("calibri.ttf", "calibrib.ttf"),
    "calibri light": ("calibril.ttf", "calibrib.ttf"),
    "cambria": ("cambria.ttc", "cambriab.ttf"),
    "bahnschrift": ("bahnschrift.ttf", "bahnschrift.ttf"),
    "century gothic": ("GOTHIC.TTF", "GOTHICB.TTF"),
    "candara": ("Candara.ttf", "Candarab.ttf"),
    "corbel": ("corbel.ttf", "corbelb.ttf"),
    "constantia": ("constan.ttf", "constanb.ttf"),
    "arial": ("arial.ttf", "arialbd.ttf"),
    "verdana": ("verdana.ttf", "verdanab.ttf"),
    "tahoma": ("tahoma.ttf", "tahomabd.ttf"),
    "trebuchet ms": ("trebuc.ttf", "trebucbd.ttf"),
    "times new roman": ("times.ttf", "timesbd.ttf"),
    "aptos": ("aptos.ttf", "aptos-bold.ttf"),
}
_BASE = 100  # px size used for measurement


@lru_cache(maxsize=64)
def _pil_font(family: str, bold: bool):
    try:
        from PIL import ImageFont
    except Exception:
        return None
    files = _FONT_FILES.get((family or "").lower().strip())
    cands = []
    if files:
        cands.append(files[1] if bold else files[0])
    cands += ["segoeuib.ttf" if bold else "segoeui.ttf", "arialbd.ttf" if bold else "arial.ttf"]
    for f in cands:
        p = _FONT_DIR / f
        if p.exists():
            try:
                return ImageFont.truetype(str(p), _BASE)
            except Exception:
                continue
    return None


@lru_cache(maxsize=64)
def _line_factor(family: str, bold: bool) -> float:
    """Single-spacing line height / em, from OS/2 win metrics (what PowerPoint uses)."""
    f = _pil_font(family, bold)
    if f is None:
        return 1.22
    try:
        from fontTools.ttLib import TTFont
        tt = TTFont(f.path, fontNumber=0, lazy=True)
        o, u = tt["OS/2"], tt["head"].unitsPerEm
        return max(1.1, min(1.45, (o.usWinAscent + o.usWinDescent) / u))
    except Exception:
        pass
    try:
        a, d = f.getmetrics()
        return max(1.15, min(1.45, (a + d) / _BASE))
    except Exception:
        return 1.22


@lru_cache(maxsize=20000)
def _word_w(word: str, family: str, bold: bool) -> float:
    """Width (inches) of `word` at 1pt."""
    f = _pil_font(family, bold)
    if f is None:
        return len(word) * 0.55 / 72
    try:
        return f.getlength(word) / _BASE / 72
    except Exception:
        return len(word) * 0.55 / 72


def text_w(text: str, family: str, size: float, bold: bool = False) -> float:
    return _word_w(text, family, bold) * size


def line_h(family: str, size: float, ls: float = 1.0, bold: bool = False) -> float:
    return size / 72 * _line_factor(family, bold) * ls


def count_lines(runs, family: str, size: float, width: float, bold_family: str = None) -> int:
    """runs: str or list of (text, bold). Greedy word wrap like PowerPoint."""
    if isinstance(runs, str):
        runs = [(runs, False)]
    width = max(width * 0.985, 0.2)       # PowerPoint wraps a hair earlier than PIL
    lines, cur = 1, 0.0
    for text, bold in runs:
        fam = (bold_family or family) if bold else family
        for piece in re.split(r"(\n)", text or ""):
            if piece == "\n":
                lines += 1
                cur = 0.0
                continue
            for tok in re.findall(r"\S+|\s+", piece):
                if tok.isspace():
                    if cur > 0:
                        cur += _word_w(" ", fam, bold) * size * len(tok)
                    continue
                w = _word_w(tok, fam, bold) * size
                if cur > 0 and cur + w > width:
                    lines += 1
                    cur = 0.0
                    # strip leading whitespace width was already added; fine
                if w > width:
                    # a single word wider than the box would be broken mid-word — never
                    # acceptable, so report it as a hopeless overflow and let fitting shrink
                    extra = int(w // width)
                    lines += extra + 50
                    cur = w - extra * width
                else:
                    cur += w
    return lines


def para_h(runs, family, size, width, ls=1.0, bold_family=None) -> float:
    bold = isinstance(runs, list) and runs and all(b for _, b in runs)
    return count_lines(runs, family, size, width, bold_family) * line_h(family, size, ls, bold)


def fit_size(runs, family, width, height, max_pt, min_pt, ls=1.0, max_lines=None, bold_family=None) -> float:
    s = max_pt
    while s > min_pt:
        n = count_lines(runs, family, s, width, bold_family)
        if (max_lines is None or n <= max_lines) and para_h(runs, family, s, width, ls, bold_family) <= height:
            return s
        s -= 1 if s > 14 else 0.5
    return min_pt


# ══════════════════════════════════════════════════════════════════════════════
#  IMAGES
# ══════════════════════════════════════════════════════════════════════════════
def prepare_image(path: str) -> tuple[Optional[str], float]:
    """Return a PPTX-safe copy (EXIF-rotated, RGB, ≤2400px, png/jpg) and its aspect."""
    try:
        from PIL import Image, ImageOps
        p = Path(path)
        if not p.exists():
            return None, 1.0
        with Image.open(p) as im:
            exif_rot = False
            try:
                exif_rot = im.getexif().get(0x0112, 1) not in (1, None)
            except Exception:
                pass
            fmt_ok = (im.format or "").upper() in ("PNG", "JPEG", "GIF", "BMP")
            big = max(im.size) > 2400
            if fmt_ok and not exif_rot and not big:
                return str(p), im.size[0] / max(im.size[1], 1)
            _CACHE_DIR.mkdir(parents=True, exist_ok=True)
            key = hashlib.md5(f"{p.resolve()}|{p.stat().st_mtime}".encode()).hexdigest()[:16]
            im2 = ImageOps.exif_transpose(im)
            if big:
                im2.thumbnail((2400, 2400))
            has_alpha = im2.mode in ("RGBA", "LA", "P")
            out = _CACHE_DIR / f"{key}.{'png' if has_alpha else 'jpg'}"
            if has_alpha:
                im2.convert("RGBA").save(out, "PNG")
            else:
                im2.convert("RGB").save(out, "JPEG", quality=90)
            return str(out), im2.size[0] / max(im2.size[1], 1)
    except Exception as e:
        print(f"[ppt_designer] image prep failed for {path}: {e}")
        return (path if Path(path).exists() else None), 1.5


@lru_cache(maxsize=256)
def _saliency_profile(path: str, axis: str) -> tuple:
    """Edge-energy profile along an axis ('x' or 'y') for smart cropping."""
    try:
        from PIL import Image, ImageFilter
        with Image.open(path) as im:
            g = im.convert("L")
            g.thumbnail((160, 160))
            e = g.filter(ImageFilter.FIND_EDGES)
            w, h = e.size
            px = e.load()
            if axis == "x":
                return tuple(sum(px[x, y] for y in range(h)) for x in range(w))
            return tuple(sum(px[x, y] for x in range(w)) for y in range(h))
    except Exception:
        return tuple()


def _best_window(profile: tuple, keep: float) -> float:
    """Start fraction of the window (length=keep fraction) with most detail, biased to centre."""
    n = len(profile)
    if n == 0 or keep >= 1:
        return (1 - keep) / 2
    k = max(1, int(round(n * keep)))
    best, best_i = -1.0, (n - k) // 2
    s = sum(profile[:k])
    for i in range(0, n - k + 1):
        if i:
            s += profile[i + k - 1] - profile[i - 1]
        centre_bias = 1 - 0.35 * abs((i + k / 2) / n - 0.5)
        if s * centre_bias > best:
            best, best_i = s * centre_bias, i
    return best_i / n


def _round_pic(pic, radius_in: float, w: float, h: float):
    try:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for gd in list(av):
            av.remove(gd)
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(min(50000, radius_in / max(min(w, h), 0.01) * 100000))}")
    except Exception:
        pass


def place_image(slide, path: str, x, y, w, h, aspect: float = None, mode: str = "cover",
                radius: float = 0.0, shadow: bool = False):
    """cover: fill box, crop with saliency. contain: fit inside box, centred."""
    if aspect is None:
        path, aspect = prepare_image(path)
    if not path:
        return None
    box_a = w / max(h, 0.01)
    if mode == "contain":
        if aspect > box_a:
            nw, nh = w, w / aspect
        else:
            nw, nh = h * aspect, h
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
        pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    else:
        pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        if abs(aspect - box_a) > 0.01:
            if aspect > box_a:            # too wide → crop sides
                keep = box_a / aspect
                start = _best_window(_saliency_profile(path, "x"), keep)
                pic.crop_left, pic.crop_right = start, max(0.0, 1 - keep - start)
            else:                          # too tall → crop top/bottom (bias upward: faces/headlines)
                keep = aspect / box_a
                start = _best_window(_saliency_profile(path, "y"), keep)
                start = min(start, (1 - keep) * 0.45)
                pic.crop_top, pic.crop_bottom = start, max(0.0, 1 - keep - start)
    if radius > 0:
        _round_pic(pic, radius, w, h)
    if shadow:
        _soft_shadow(pic)
    return pic


def justified_rows(aspects: list[float], w: float, h: float, gap: float, cap_h: float = 0.0):
    """Google-Photos-style justified layout. Returns list of (x, y, w, h) relative to box, no crops.
    cap_h reserves caption space under every row (captions don't scale with the images)."""
    n = len(aspects)
    if cap_h:
        best = None
        for rows in range(1, n + 1):
            boxes = _justified_fixed(aspects, w, h - rows * cap_h, gap, rows)
            if boxes is None:
                continue
            area = sum(b[2] * b[3] for b in boxes)
            if best is None or area > best[0]:
                best = (area, boxes, rows)
        if best:
            _, boxes, rows = best
            ys = sorted({round(b[1], 4) for b in boxes})
            used = max(b[1] + b[3] for b in boxes) + rows * cap_h
            off = 0.0                      # top-aligned: sits right under its heading
            return [(bx, by + ys.index(round(by, 4)) * cap_h + off - min(ys), bw, bh) for bx, by, bw, bh in boxes]
    best = None
    for rows in range(1, n + 1):
        # balanced sequential partition by aspect sum
        target = sum(aspects) / rows
        parts, cur, acc = [], [], 0.0
        for i, a in enumerate(aspects):
            cur.append(i)
            acc += a
            remaining_rows = rows - len(parts) - 1
            if acc >= target * 0.92 and remaining_rows > 0 and (n - i - 1) >= remaining_rows:
                parts.append(cur)
                cur, acc = [], 0.0
        if cur:
            parts.append(cur)
        heights = [(w - gap * (len(p) - 1)) / sum(aspects[j] for j in p) for p in parts]
        total = sum(heights) + gap * (len(parts) - 1)
        scale = min(1.0, h / total)
        score = -(w * scale) * (total * scale)          # maximise covered area
        if best is None or score < best[0]:
            best = (score, parts, heights, scale)
    _, parts, heights, scale = best
    boxes = [None] * n
    total_h = (sum(heights) + gap * (len(parts) - 1)) * scale
    yy = (h - total_h) / 2
    for p, rh in zip(parts, heights):
        rh *= scale
        row_w = sum(aspects[j] * rh for j in p) + gap * (len(p) - 1) * scale
        xx = (w - row_w) / 2
        for j in p:
            bw = aspects[j] * rh
            boxes[j] = (xx, yy, bw, rh)
            xx += bw + gap * scale
        yy += rh + gap * scale
    return boxes


def _justified_fixed(aspects, w, h, gap, rows):
    """Justified layout with an exact number of rows (top-aligned). None if impossible."""
    n = len(aspects)
    if rows > n or h <= 0.3:
        return None
    target = sum(aspects) / rows
    parts, cur, acc = [], [], 0.0
    for i, a in enumerate(aspects):
        cur.append(i)
        acc += a
        left_rows = rows - len(parts) - 1
        if (acc >= target * 0.92 and left_rows > 0 and (n - i - 1) >= left_rows) or (n - i - 1) == left_rows > 0:
            parts.append(cur)
            cur, acc = [], 0.0
    if cur:
        parts.append(cur)
    if len(parts) != rows:
        return None
    heights = [(w - gap * (len(p) - 1)) / sum(aspects[j] for j in p) for p in parts]
    total = sum(heights) + gap * (rows - 1)
    scale = min(1.0, h / total)
    boxes = [None] * n
    yy = 0.0
    for p, rh in zip(parts, heights):
        rh *= scale
        row_w = sum(aspects[j] * rh for j in p) + gap * (len(p) - 1)
        xx = (w - row_w) / 2
        for j in p:
            bw = aspects[j] * rh
            boxes[j] = (xx, yy, bw, rh)
            xx += bw + gap
        yy += rh + gap
    return boxes


# ══════════════════════════════════════════════════════════════════════════════
#  LOW-LEVEL DRAWING
# ══════════════════════════════════════════════════════════════════════════════
def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#")[:6].upper())


def _set_alpha(fill_fore_color_el, alpha: float):
    clr = fill_fore_color_el
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def _soft_shadow(shape, alpha: float = 0.14, blur_pt: float = 18, dist_pt: float = 4):
    try:
        spPr = shape._element.spPr
        for old in spPr.findall(qn("a:effectLst")):
            spPr.remove(old)
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"))
        sh.set("blurRad", str(int(blur_pt * 12700)))
        sh.set("dist", str(int(dist_pt * 12700)))
        sh.set("dir", "5400000")
        sh.set("algn", "t")
        sh.set("rotWithShape", "0")
        c = etree.SubElement(sh, qn("a:srgbClr"))
        c.set("val", "0F172A")
        a = etree.SubElement(c, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))
    except Exception:
        pass


def box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, radius=0.0, alpha=None,
        shape=None, shadow=False):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / max(min(w, h), 0.01))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None and alpha < 1:
            _set_alpha(shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if shadow:
        _soft_shadow(shp)
    shp.text_frame.text = ""
    return shp


def gradient_box(slide, x, y, w, h, color: str, a_from: float, a_to: float, angle: int = 0):
    """Rectangle with a single-colour alpha gradient (angle in degrees, 0 = left→right)."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    shp.line.fill.background()
    shp.shadow.inherit = False
    spPr = shp._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    grad = etree.Element(qn("a:gradFill"))
    grad.set("rotWithShape", "1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, a in ((0, a_from), (100000, a_to)):
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(pos))
        c = etree.SubElement(gs, qn("a:srgbClr"))
        c.set("val", color)
        al = etree.SubElement(c, qn("a:alpha"))
        al.set("val", str(int(a * 100000)))
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(int(angle * 60000)))
    lin.set("scaled", "0")
    geom = spPr.find(qn("a:prstGeom"))
    geom.addnext(grad)
    return shp


def line_shape(slide, x1, y1, x2, y2, color, width_pt=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(width_pt)
    return c


def text(slide, x, y, w, h, paras, *, font, size, color, bold=False, align="l", anchor="t",
         ls=1.0, space_after=0.0, italic=False, spacing=None):
    """paras: list of str | list[(text, opts)]. opts keys: bold, color, size, font, italic."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = ls
        if space_after and i < len(paras) - 1:
            p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = _rgb(o.get("color", color))
            if spacing is not None or o.get("spacing") is not None:
                r.font._element.set("spc", str(int((o.get("spacing", spacing) or 0) * 100)))
    return tb


# ══════════════════════════════════════════════════════════════════════════════
#  DENSITY PROFILES
# ══════════════════════════════════════════════════════════════════════════════
# Fitting always picks the LARGEST size that fits, so density mostly sets how
# small text may get (dense decks carry more words) and a gentler ceiling.
DENSITY = {
    "light":    dict(title_max=40, title_min=28, body_max=24, body_min=15, head_max=22, lead_max=22),
    "balanced": dict(title_max=38, title_min=26, body_max=22, body_min=13, head_max=20, lead_max=20),
    "dense":    dict(title_max=36, title_min=24, body_max=20, body_min=11, head_max=18, lead_max=18),
}


def _items(sp: dict) -> list[dict]:
    """Unify bullets / steps / cards into [{head, text, date}]."""
    out = []
    for key in ("bullets", "steps", "cards", "items"):
        for b in sp.get(key) or []:
            if isinstance(b, str):
                if b.strip():
                    out.append({"head": "", "text": b.strip()})
            elif isinstance(b, dict):
                head = str(b.get("head") or b.get("bold") or b.get("header") or b.get("title") or "").strip()
                txt = str(b.get("text") or b.get("desc") or b.get("description") or "").strip()
                date = str(b.get("date") or b.get("when") or "").strip()
                if head or txt:
                    out.append({"head": head, "text": txt, "date": date})
        if out:
            return out
    return out


def _runs_for(item, head_font, body_font) -> list:
    if item.get("head") and item.get("text"):
        return [(item["head"] + "  ", True), (item["text"], False)]
    return [(item.get("head") or item.get("text") or "", bool(item.get("head")))]


# ══════════════════════════════════════════════════════════════════════════════
#  RENDERER
# ══════════════════════════════════════════════════════════════════════════════
def rich(s: str) -> list[tuple[str, bool]]:
    """'Scores **6 evidence sources** …' → [('Scores ', False), ('6 evidence sources', True), …]."""
    out = []
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", s or "")):
        if part:
            out.append((part, i % 2 == 1))
    return out or [("", False)]


def plain(s: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", s or "")


# ── icons: Windows 11 "Segoe Fluent Icons" glyphs rendered to PNG (portable in any viewer) ──
_ICON_FONT = _FONT_DIR / "SegoeIcons.ttf"
_ICONS = [  # (keyword regex, codepoint) — first match wins
    (r"\b(ai|ml|llm|genai|gpt|robot|bot|model|neural|intelligen)", 0xE99A),
    (r"\b(search|retriev|find|discover|semantic|faiss|bm25|index)", 0xE721),
    (r"\b(dashboard|analytic|chart|metric|report|trend|heatmap|insight|visuali)", 0xE9D2),
    (r"\b(assess|quiz|mcq|test|exam|question|evaluat|grad)", 0xEADF),
    (r"\b(learn|course|training|educat|skill|upskill|capacity|curricul|study)", 0xE7BE),
    (r"\b(multilingual|language|translat|hindi|regional|bilingual)", 0xE8C1),
    (r"\b(voice|speech|mic|audio|speak)", 0xE720),
    (r"\b(chat|assistant|conversation|message|support|query)", 0xE8F2),
    (r"\b(secur|privacy|complian|protect|safe|trust|dpdp|regulat|audit)", 0xEA18),
    (r"\b(lock|encrypt|auth|password|login|access)", 0xE72E),
    (r"\b(cloud|server|deploy|host|saas|intranet|infra)", 0xE753),
    (r"\b(data|database|storage|postgres|sql|record)", 0xEE94),
    (r"\b(sync|integrat|connect|api|interop|crosswalk|adapter)", 0xE895),
    (r"\b(fast|speed|real-?time|instant|quick|live|latency|performance)", 0xE945),
    (r"\b(people|user|team|stakeholder|employee|official|citizen|farmer|student|learner|workforce|community)", 0xE716),
    (r"\b(point|reward|gamif|badge|streak|award|star|rating)", 0xE734),
    (r"\b(cost|price|econom|revenue|money|budget|fund|roi|profit|licen)", 0xE8C7),
    (r"\b(document|pdf|ocr|scan|file|ppt|paper|text)", 0xE8A5),
    (r"\b(global|world|national|india|region|district|map|location)", 0xE774),
    (r"\b(idea|innovat|novel|unique|creative|vision)", 0xEA80),
    (r"\b(goal|target|mission|objective|aim|gap)", 0xE7C1),
    (r"\b(time|schedule|calendar|month|year|deadline|plan)", 0xE787),
    (r"\b(health|medical|care|wellness|patient)", 0xE9D9),
    (r"\b(mobile|phone|app)\b", 0xE8EA),
    (r"\b(laptop|device|hardware|sensor|iot|chip|edge)", 0xE950),
    (r"\b(pipeline|engine|process|workflow|automat|system|architecture)", 0xE9F5),
    (r"\b(tool|build|develop|code|implement|stack|framework)", 0xEC7A),
    (r"\b(check|verif|valid|accura|confidence|proof|quality|reliab)", 0xE73E),
    (r"\b(risk|challenge|warning|problem|issue|barrier)", 0xE7BA),
    (r"\b(benefit|impact|value|improv|growth|scale)", 0xE8E1),
    (r"\b(network|mesh|graph|node)", 0xEBD2),
    (r"\b(library|resource|reference|research|source|catalog)", 0xE8F1),
    (r"\b(setting|config|custom)", 0xE713),
]
_ICON_FALLBACK = [0xEA80, 0xE734, 0xE9F5, 0xE8E1, 0xE7C1, 0xE73E]


def icon_for(text_: str, k: int = 0) -> int:
    low = (text_ or "").lower()
    for rx, cp in _ICONS:
        if re.search(rx, low):
            return cp
    return _ICON_FALLBACK[k % len(_ICON_FALLBACK)]


@lru_cache(maxsize=256)
def icon_png(codepoint: int, color: str, px: int = 128) -> Optional[str]:
    try:
        from PIL import Image, ImageDraw, ImageFont
        if not _ICON_FONT.exists():
            return None
        _CACHE_DIR.mkdir(parents=True, exist_ok=True)
        out = _CACHE_DIR / f"icon_{codepoint:X}_{color}_{px}.png"
        if not out.exists():
            f = ImageFont.truetype(str(_ICON_FONT), int(px * 0.8))
            im = Image.new("RGBA", (px, px), (0, 0, 0, 0))
            d = ImageDraw.Draw(im)
            l, t, r, b = d.textbbox((0, 0), chr(codepoint), font=f)
            rgb = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
            d.text(((px - (r - l)) / 2 - l, (px - (b - t)) / 2 - t), chr(codepoint), font=f, fill=rgb + (255,))
            im.save(out)
        return str(out)
    except Exception as e:
        print(f"[ppt_designer] icon failed: {e}")
        return None


class DeckRenderer:
    def __init__(self, deck: dict, out_path: str):
        self.deck = deck
        self.out = out_path
        self.t = deck["theme"] if deck.get("theme", {}).get("series") else _finish_theme(deck["theme"])
        self.d = DENSITY.get(deck.get("density", "light"), DENSITY["light"])
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(SW), E(SH)
        self._blank = self.prs.slide_layouts[6]
        self._img_cache: dict[str, tuple] = {}
        self._side = 1              # alternates image side (1 = right)
        self._prev_variant = ""
        self.hf = self.t["head_font"]
        self.bf = self.t["body_font"]
        self.bf_bold = "Segoe UI Semibold" if self.bf.lower().startswith("segoe") else self.bf
        self.logo = deck.get("logo")

    # ── helpers ───────────────────────────────────────────────────────────────
    def _img(self, path):
        if path not in self._img_cache:
            self._img_cache[path] = prepare_image(path)
        return self._img_cache[path]

    def _images(self, sp) -> list[tuple[str, float]]:
        out = []
        for p in sp.get("images") or []:
            pp, a = self._img(p)
            if pp:
                out.append((pp, a))
        return out

    def _bg(self, s, color=None):
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color or self.t["bg"])

    def _footer(self, s, idx, total):
        t = self.t
        y = SH - 0.45
        label = (self.deck.get("footer") or self.deck.get("title") or "")[:60]
        if label:
            text(s, MX, y, 6, 0.25, label, font=self.bf, size=9, color=t["muted"])
        text(s, SW - MX - 1.2, y, 1.2, 0.25, f"{idx:02d}", font=self.bf_bold, size=9,
             color=t["muted"], align="r")
        if self.logo:
            p, a = self._img(self.logo)
            if p:
                hh = 0.32
                place_image(s, p, SW - MX - 1.35 - hh * a, y - 0.05, hh * a, hh, a, mode="contain")

    def _header(self, s, sp, x=MX, w=SW - 2 * MX, y=MT, kicker=None, max_lines=2):
        """Accent tick + (kicker) + fitted title + (subtitle). Returns bottom y."""
        t, d = self.t, self.d
        title = sp.get("title") or ""
        box(s, x, y, 0.55, 0.07, fill=t["accent"])
        yy = y + 0.22
        if kicker:
            text(s, x + 0.7, y - 0.07, w - 0.7, 0.25, kicker.upper(), font=self.bf_bold, size=10,
                 color=t["accent_text"], spacing=2)
        size = fit_size(title, self.hf, w, line_h(self.hf, d["title_max"]) * max_lines + 0.05,
                        d["title_max"], d["title_min"] - 4, max_lines=max_lines)
        th = para_h(title, self.hf, size, w)
        text(s, x, yy, w, th + 0.05, title, font=self.hf, size=size, color=t["text"])
        yy += th
        sub = sp.get("subtitle") or ""
        if sub:
            ss = fit_size(sub, self.bf, w, 0.8, 18, 13, max_lines=2)
            shh = para_h(sub, self.bf, ss, w)
            text(s, x, yy + 0.1, w, shh + 0.05, sub, font=self.bf, size=ss, color=t["muted"])
            yy += 0.1 + shh
        return yy + 0.38

    def _notes(self, s, sp):
        if sp.get("notes"):
            try:
                s.notes_slide.notes_text_frame.text = str(sp["notes"])
            except Exception:
                pass

    # ── item list (the workhorse) ─────────────────────────────────────────────
    def _list_fit(self, items, w, h, style):
        """Largest body size so all items fit in (w,h). style: 'stack'|'inline'."""
        d = self.d
        n = max(len(items), 1)
        for size in _sizes(d["body_max"], max(d["body_min"] - 2, 10)):
            gap = max(0.14, size / 72 * 0.9)
            total = 0.0
            for it in items:
                total += self._item_h(it, size, w, style)
            total += gap * (n - 1)
            if total <= h:
                return size, gap, total
        size = max(d["body_min"] - 2, 10)
        gap = 0.12
        return size, gap, sum(self._item_h(it, size, w, style) for it in items) + gap * (n - 1)

    def _item_h(self, it, size, w, style):
        if style == "stack" and it.get("head") and it.get("text"):
            hs = size + 2
            return (para_h(it["head"], self.bf_bold, hs, w, bold_family=self.bf_bold)
                    + 0.04 + para_h(it["text"], self.bf, size, w, 1.05))
        return para_h(_runs_for(it, self.hf, self.bf), self.bf, size, w, 1.05, bold_family=self.bf_bold)

    def _draw_item(self, s, it, x, y, w, size, style, color=None):
        t = self.t
        color = color or t["text"]
        if style == "stack" and it.get("head") and it.get("text"):
            hs = size + 2
            hh = para_h(it["head"], self.bf_bold, hs, w, bold_family=self.bf_bold)
            text(s, x, y, w, hh + 0.02, it["head"], font=self.bf_bold, size=hs, color=color)
            bh = para_h(it["text"], self.bf, size, w, 1.05)
            text(s, x, y + hh + 0.04, w, bh + 0.02, it["text"], font=self.bf, size=size,
                 color=t["muted"], ls=1.05)
            return hh + 0.04 + bh
        runs = _runs_for(it, self.hf, self.bf)
        paras = [[(r, {"bold": b, "font": self.bf_bold if b else self.bf,
                       "color": color if b or not it.get("head") else t["muted"]}) for r, b in runs]]
        hh = para_h(runs, self.bf, size, w, 1.05, bold_family=self.bf_bold)
        text(s, x, y, w, hh + 0.02, paras, font=self.bf, size=size, color=color, ls=1.05)
        return hh

    def _list(self, s, items, x, y, w, h, marker="dot", valign="m"):
        """Items with markers, fitted and vertically balanced."""
        t = self.t
        style = "stack" if sum(1 for i in items if i.get("head") and i.get("text")) >= max(1, len(items) * 0.6) \
            and len(items) <= 5 else "inline"
        mk_w = 0.55 if marker == "num" else 0.36
        tw = w - mk_w
        size, gap, total = self._list_fit(items, tw, h, style)
        yy = y + (max(0.0, (h - total) / 2) if valign == "m" else 0)
        if valign == "m":
            yy = _vpos(y, h, total)
        for i, it in enumerate(items):
            ih = self._item_h(it, size, tw, style)
            first_line = line_h(self.bf_bold, size + (2 if style == "stack" and it.get("head") and it.get("text") else 0))
            if marker == "num":
                d = min(0.42, max(0.3, first_line * 1.05))
                box(s, x, yy + (first_line - d) / 2, d, d, fill=t["accent"], shape=MSO_SHAPE.OVAL)
                text(s, x, yy + (first_line - d) / 2, d, d, str(i + 1), font=self.bf_bold,
                     size=max(10, d * 72 * 0.42), color=t["on_accent"], align="c", anchor="m")
            elif marker == "bar":
                box(s, x, yy + 0.04, 0.06, max(0.2, ih - 0.08), fill=t["accent"] if i % 2 == 0 else t["accent2"])
            else:
                dd = max(0.1, size / 72 * 0.5)
                box(s, x + 0.04, yy + first_line / 2 - dd / 2, dd, dd, fill=t["accent"], shape=MSO_SHAPE.OVAL)
            self._draw_item(s, it, x + mk_w, yy, tw, size, style)
            yy += ih + gap
        return size

    def _list_grid(self, s, items, x, y, w, h, cols=2):
        """Row-aligned multi-column list: item i sits in row i//cols, so rows line up."""
        t = self.t
        g = 0.6
        cw = (w - g * (cols - 1)) / cols
        mk = 0.36
        tw = cw - mk
        rows = [items[i:i + cols] for i in range(0, len(items), cols)]
        size, gap, row_h = self.d["body_min"], 0.2, []
        for sz in _sizes(self.d["body_max"], max(self.d["body_min"] - 2, 10)):
            gap = max(0.2, sz / 72 * 1.1)
            row_h = [max(self._item_h(it, sz, tw, "inline") for it in r) for r in rows]
            if sum(row_h) + gap * (len(rows) - 1) <= h:
                size = sz
                break
        total = sum(row_h) + gap * (len(rows) - 1)
        yy = _vpos(y, h, total)
        for r, rh in zip(rows, row_h):
            for c, it in enumerate(r):
                cx = x + c * (cw + g)
                dd = max(0.1, size / 72 * 0.5)
                fl = line_h(self.bf, size, 1.05)
                box(s, cx + 0.04, yy + fl / 2 - dd / 2, dd, dd, fill=t["accent"], shape=MSO_SHAPE.OVAL)
                self._draw_item(s, it, cx + mk, yy, tw, size, "inline")
            yy += rh + gap

    def _cards(self, s, items, x, y, w, h, numbered=True):
        """Grid of rounded cards, uniform font size, equal heights per row."""
        t = self.t
        n = len(items)
        words = sum(len((i.get("head", "") + " " + i.get("text", "")).split()) for i in items) / max(n, 1)
        if n <= 3:
            cols = n
        elif n == 4:
            cols = 4 if words < 16 and w > 10 else 2
        elif n in (5, 6):
            cols = 3
        else:
            cols = 4
        cols = max(1, min(cols, n))
        rows = math.ceil(n / cols)
        base, rem = divmod(n, rows)                       # balanced rows (5 → 3+2), each stretched full width
        counts = [base + (1 if k < rem else 0) for k in range(rows)]
        g = 0.25
        cw = (w - g * (max(counts) - 1)) / max(counts)    # narrowest card decides the font size
        ch_max = (h - g * (rows - 1)) / rows
        pad = 0.28
        tw = cw - 2 * pad
        badge = 0.42 if numbered else 0.0
        top_extra = (badge + 0.18) if numbered else 0.12
        d = self.d
        size = d["body_min"] - 1
        for sz in _sizes(d["body_max"] - 1, max(d["body_min"] - 2, 10)):
            ok = True
            for it in items:
                need = top_extra + pad * 2 + self._card_text_h(it, sz, tw)
                if need > ch_max:
                    ok = False
                    break
            if ok:
                size = sz
                break
        heights = [top_extra + pad * 2 + self._card_text_h(it, size, tw) for it in items]
        ch = min(ch_max, max(max(heights) * 1.45, 1.6))   # cards grow into the free space instead of leaving it
        total_h = rows * ch + (rows - 1) * g
        y0 = _vpos(y, h, total_h)
        slots = [(r_, c_, counts[r_]) for r_ in range(rows) for c_ in range(counts[r_])]
        for i, it in enumerate(items):
            r, c, in_row = slots[i]
            cw_r = (w - g * (in_row - 1)) / in_row
            cx = x + c * (cw_r + g)
            cy = y0 + r * (ch + g)
            cw_i = cw_r
            box(s, cx, cy, cw_i, ch, fill=t["surface"], line=None if not t["dark"] else t["line"],
                radius=0.14, shadow=not t["dark"])
            row_items = [items[j] for j in range(n) if slots[j][0] == r]
            nat = max(top_extra + pad * 2 + self._card_text_h(o, size, cw_i - 2 * pad) for o in row_items)
            ty = cy + pad + max(0.0, (ch - nat) / 2)          # one shared offset per row → badges line up
            acc = t["series"][i % 3]
            if numbered:
                box(s, cx + pad, ty, badge, badge, fill=acc, radius=0.1)
                text(s, cx + pad, ty, badge, badge, f"{i + 1:02d}", font=self.bf_bold, size=12,
                     color=on_color(acc), align="c", anchor="m")
                ty += badge + 0.18
            else:
                box(s, cx + pad, ty, 0.5, 0.06, fill=acc)
                ty += 0.2
            self._card_text(s, it, cx + pad, ty, cw_i - 2 * pad, size)
        return size

    def _card_text_h(self, it, size, w):
        hs = size + 3
        h = 0.0
        if it.get("head"):
            h += para_h(it["head"], self.bf_bold, hs, w, bold_family=self.bf_bold) + 0.08
        if it.get("text"):
            h += para_h(it["text"], self.bf, size, w, 1.05)
        return h

    def _card_text(self, s, it, x, y, w, size, color=None):
        t = self.t
        hs = size + 3
        if it.get("head"):
            hh = para_h(it["head"], self.bf_bold, hs, w, bold_family=self.bf_bold)
            text(s, x, y, w, hh + 0.02, it["head"], font=self.bf_bold, size=hs, color=color or t["text"])
            y += hh + 0.08
        if it.get("text"):
            bh = para_h(it["text"], self.bf, size, w, 1.05)
            text(s, x, y, w, bh + 0.02, it["text"], font=self.bf, size=size,
                 color=t["muted"] if not color else color, ls=1.05)

    # ── composite (reference-grade infographic) slides → ppt_composer ─────────
    def _s_sections(self, s, sp, idx, total):
        from app.services.ppt_composer import render_sections
        render_sections(self, s, sp, idx, total)

    def _s_title_sections(self, s, sp, idx, total):
        from app.services.ppt_composer import render_title_sections
        render_title_sections(self, s, sp, idx, total)

    # ── slide builders ────────────────────────────────────────────────────────
    def render(self):
        slides = self.deck.get("slides", [])
        total = len(slides)
        for i, sp in enumerate(slides):
            kind = sp.get("kind") or "content"
            if kind == "title" and sp.get("sections"):
                kind = "title_sections"
            elif kind != "sections":
                sp = _deep_plain(sp)          # **bold** markup is only rendered by the composite layout
            fn = getattr(self, f"_s_{kind}", None) or self._s_content
            if kind not in ("content", "cards"):
                self._prev_variant = kind
            s = self.prs.slides.add_slide(self._blank)
            self._bg(s)
            try:
                fn(s, sp, i + 1, total)
            except Exception as e:                       # never lose a slide
                print(f"[ppt_designer] slide {i + 1} ({kind}) failed: {e}; using content layout")
                for shp in list(s.shapes):
                    shp._element.getparent().remove(shp._element)
                try:
                    self._s_content(s, _as_plain_content(sp), i + 1, total)
                except Exception as e2:
                    text(s, MX, MT, SW - 2 * MX, 1, sp.get("title", ""), font=self.hf, size=32, color=self.t["text"])
                    print(f"[ppt_designer] fallback failed too: {e2}")
            self._notes(s, sp)
            yield f"  ✓ Slide {i + 1}/{total}: {sp.get('title', '')[:60]}"
        self.prs.core_properties.title = self.deck.get("title", "")[:200]
        self.prs.save(self.out)

    # ---- title ---------------------------------------------------------------
    def _s_title(self, s, sp, idx, total):
        t = self.t
        imgs = self._images(sp)
        title = sp.get("title") or self.deck.get("title", "")
        sub = sp.get("subtitle") or ""
        meta = sp.get("body") or ""
        if imgs and imgs[0][1] >= 1.2:
            p, a = imgs[0]
            place_image(s, p, 0, 0, SW, SH, a, mode="cover")
            gradient_box(s, 0, 0, SW, SH, "07070B", 0.88, 0.15, angle=0)
            gradient_box(s, 0, SH * 0.45, SW, SH * 0.55, "07070B", 0.0, 0.55, angle=90)
            self._title_block(s, title, sub, meta, MX + 0.1, 0, SW * 0.58, SH,
                              tc="FFFFFF", mc="D4D4D8", ac=t["accent"])
            self._title_logo(s)
            return
        if imgs:
            p, a = imgs[0]
            pw = min(max(SH * a, SW * 0.38), SW * 0.5)
            place_image(s, p, SW - pw, 0, pw, SH, a, mode="cover")
            self._title_block(s, title, sub, meta, MX + 0.1, 0, SW - pw - MX - 0.6, SH)
            return
        # no image: bold geometry, variant from theme darkness
        if t["dark"]:
            box(s, SW - 5.2, -1.6, 7.0, 7.0, fill=t["accent"], alpha=0.10, shape=MSO_SHAPE.OVAL)
            box(s, SW - 3.6, 3.4, 5.0, 5.0, fill=t["accent2"], alpha=0.10, shape=MSO_SHAPE.OVAL)
            box(s, SW - 4.4, 1.2, 1.1, 1.1, fill=t["accent3"], alpha=0.85, shape=MSO_SHAPE.OVAL)
        else:
            box(s, SW - 4.6, 0, 4.6, SH, fill=t["surface2"])
            box(s, SW - 4.6, 0, 0.12, SH, fill=t["accent"])
            box(s, SW - 3.4, 1.4, 2.2, 2.2, fill=t["accent"], alpha=0.9, shape=MSO_SHAPE.OVAL)
            box(s, SW - 2.3, 3.2, 1.5, 1.5, fill=t["accent2"], alpha=0.85, shape=MSO_SHAPE.OVAL)
            box(s, SW - 3.9, 4.4, 0.8, 0.8, fill=t["accent3"], alpha=0.9, shape=MSO_SHAPE.OVAL)
        self._title_logo(s)
        self._title_block(s, title, sub, meta, MX + 0.1, 0, SW - 5.6 - MX, SH)

    def _title_logo(self, s):
        if self.logo:
            p, a = self._img(self.logo)
            if p:
                hh = 0.55 if a < 2.5 else 0.45
                place_image(s, p, MX + 0.1, 0.5, min(hh * a, 3.0), hh, a, mode="contain")

    def _title_block(self, s, title, sub, meta, x, y, w, h, tc=None, mc=None, ac=None):
        t = self.t
        tc, mc, ac = tc or t["text"], mc or t["muted"], ac or t["accent"]
        ts = fit_size(title, self.hf, w, 2.9, 60, 30, max_lines=3)
        th = para_h(title, self.hf, ts, w)
        ss = fit_size(sub, self.bf, w, 1.1, 22, 14, max_lines=3) if sub else 0
        sh_ = para_h(sub, self.bf, ss, w) if sub else 0
        mh = 0.35 if meta else 0
        block = 0.3 + th + (0.25 + sh_ if sub else 0) + (0.35 + mh if meta else 0)
        yy = y + (h - block) / 2
        box(s, x, yy, 0.8, 0.09, fill=ac)
        yy += 0.3
        text(s, x, yy, w, th + 0.05, title, font=self.hf, size=ts, color=tc)
        yy += th + 0.25
        if sub:
            text(s, x, yy, w, sh_ + 0.05, sub, font=self.bf, size=ss, color=mc)
            yy += sh_ + 0.35
        if meta:
            text(s, x, yy, w, mh, meta, font=self.bf_bold, size=12, color=mc, spacing=1)

    # ---- section -------------------------------------------------------------
    def _s_section(self, s, sp, idx, total):
        t = self.t
        self._bg(s, t["surface2"] if not t["dark"] else t["surface"])
        num = sp.get("number") or f"{sp.get('_section_no', idx):02d}"
        box(s, 0, 0, 0.18, SH, fill=t["accent"])
        text(s, MX + 0.2, 1.3, 4, 1.6, str(num), font=self.hf, size=96, color=t["accent_text"])
        w = SW - 2 * MX - 0.4
        ts = fit_size(sp.get("title", ""), self.hf, w, 1.9, 48, 30, max_lines=2)
        th = para_h(sp.get("title", ""), self.hf, ts, w)
        text(s, MX + 0.2, 3.2, w, th + 0.05, sp.get("title", ""), font=self.hf, size=ts, color=t["text"])
        if sp.get("subtitle") or sp.get("body"):
            sub = sp.get("subtitle") or sp.get("body")
            text(s, MX + 0.2, 3.3 + th, w * 0.75, 1.2, sub, font=self.bf, size=18, color=t["muted"])

    # ---- closing -------------------------------------------------------------
    def _s_closing(self, s, sp, idx, total):
        t = self.t
        imgs = self._images(sp)
        if imgs:
            p, a = imgs[0]
            place_image(s, p, 0, 0, SW, SH, a, mode="cover")
            gradient_box(s, 0, 0, SW, SH, "07070B", 0.75, 0.75)
            tc, mc = "FFFFFF", "D4D4D8"
        else:
            self._bg(s, t["accent"] if not t["dark"] else t["bg"])
            tc = t["on_accent"] if not t["dark"] else t["text"]
            mc = _mix(tc, t["accent"] if not t["dark"] else t["bg"], 0.25)
            if t["dark"]:
                box(s, -2.0, -2.5, 6.5, 6.5, fill=t["accent"], alpha=0.12, shape=MSO_SHAPE.OVAL)
                box(s, SW - 3.5, SH - 3.0, 5.5, 5.5, fill=t["accent2"], alpha=0.12, shape=MSO_SHAPE.OVAL)
            else:
                box(s, -2.0, -2.5, 6.5, 6.5, fill="FFFFFF", alpha=0.08, shape=MSO_SHAPE.OVAL)
                box(s, SW - 3.5, SH - 3.0, 5.5, 5.5, fill="FFFFFF", alpha=0.08, shape=MSO_SHAPE.OVAL)
        w = SW - 4
        title = sp.get("title") or "Thank you"
        ts = fit_size(title, self.hf, w, 2.2, 66, 32, max_lines=2)
        th = para_h(title, self.hf, ts, w)
        items = _items(sp)
        sub = sp.get("subtitle") or sp.get("body") or ""
        extra = " · ".join(filter(None, [(i.get("head") + ": " + i["text"]) if i.get("head") and i.get("text")
                                          else (i.get("head") or i.get("text")) for i in items]))
        sub_h = para_h(sub, self.bf, 20, w) if sub else 0
        ex_h = para_h(extra, self.bf, 14, w) if extra else 0
        block = th + (0.3 + sub_h if sub else 0) + (0.35 + ex_h if extra else 0)
        yy = (SH - block) / 2
        text(s, 2, yy, w, th + 0.05, title, font=self.hf, size=ts, color=tc, align="c")
        yy += th + 0.3
        if sub:
            text(s, 2, yy, w, sub_h + 0.05, sub, font=self.bf, size=20, color=mc, align="c")
            yy += sub_h + 0.35
        if extra:
            text(s, 2, yy, w, ex_h + 0.05, extra, font=self.bf, size=14, color=mc, align="c")

    # ---- quote ---------------------------------------------------------------
    def _s_quote(self, s, sp, idx, total):
        t = self.t
        q = sp.get("quote") or {}
        qt = (q.get("text") if isinstance(q, dict) else str(q)) or sp.get("body") or sp.get("title", "")
        author = (q.get("author") if isinstance(q, dict) else "") or sp.get("subtitle", "")
        imgs = self._images(sp)
        x, w = MX + 0.9, SW - 2 * MX - 1.8
        if imgs:
            p, a = imgs[0]
            pw = min(max(SH * a, SW * 0.3), SW * 0.42)
            place_image(s, p, SW - pw, 0, pw, SH, a)
            w = SW - pw - x - 0.7
        text(s, x - 0.35, 0.7, 2, 2, "\u201C", font="Georgia", size=160, color=t["accent_text"])
        qs = fit_size(qt, "Georgia", w, 3.6, 40, 20, ls=1.1)
        qh = para_h(qt, "Georgia", qs, w, 1.1)
        yy = max(2.0, (SH - qh - 0.8) / 2 + 0.2)
        text(s, x, yy, w, qh + 0.1, qt, font="Georgia", size=qs, color=t["text"], ls=1.1, italic=True)
        if author:
            box(s, x, yy + qh + 0.35, 0.5, 0.05, fill=t["accent"])
            text(s, x + 0.7, yy + qh + 0.22, w - 0.7, 0.4, author, font=self.bf_bold, size=15, color=t["muted"])
        if sp.get("title") and sp.get("title") != qt:
            text(s, x, 0.6, w, 0.4, sp["title"].upper(), font=self.bf_bold, size=11,
                 color=t["muted"], spacing=2)
        self._footer(s, idx, total)

    # ---- agenda --------------------------------------------------------------
    def _s_agenda(self, s, sp, idx, total):
        t = self.t
        items = _items(sp)
        if not items:
            return self._s_content(s, sp, idx, total)
        left_w = 4.2
        box(s, 0, 0, left_w, SH, fill=t["surface2"] if not t["dark"] else t["surface"])
        box(s, 0, 0, 0.12, SH, fill=t["accent"])
        ts = fit_size(sp.get("title", "Agenda"), self.hf, left_w - 1.2, 2.5, 44, 26, max_lines=3)
        th = para_h(sp.get("title", "Agenda"), self.hf, ts, left_w - 1.2)
        text(s, MX, (SH - th) / 2, left_w - 1.2, th + 0.1, sp.get("title", "Agenda"), font=self.hf, size=ts, color=t["text"])
        x0 = left_w + 0.8
        w = SW - x0 - MX
        cols = 2 if len(items) > 5 else 1
        per = math.ceil(len(items) / cols)
        cw = (w - GAP * (cols - 1)) / cols
        row_h = min(1.0, (SH - 2 * MT - 0.4) / per)
        top = (SH - row_h * per) / 2
        size = min(fit_size(max((i.get("head") or i.get("text") for i in items), key=len),
                            self.bf_bold, cw - 1.0, row_h * 0.85, 24, 13, max_lines=2), 24)
        for i, it in enumerate(items):
            c, r = divmod(i, per)
            cx, cy = x0 + c * (cw + GAP), top + r * row_h
            text(s, cx, cy, 0.9, row_h, f"{i + 1:02d}", font=self.hf, size=size + 4,
                 color=t["accent_text"], anchor="m")
            label = it.get("head") or it.get("text")
            text(s, cx + 0.95, cy, cw - 0.95, row_h, label, font=self.bf_bold, size=size,
                 color=t["text"], anchor="m")
            if r < per - 1 and (c * per + r + 1) < len(items):
                line_shape(s, cx + 0.95, cy + row_h, cx + cw, cy + row_h, t["line"], 0.75)
        self._footer(s, idx, total)

    # ---- content (bullets / body / image) -------------------------------------
    def _s_content(self, s, sp, idx, total):
        t = self.t
        items = _items(sp)
        body = (sp.get("body") or "").strip()
        imgs = self._images(sp)
        if not items and not body and not imgs:
            return self._s_section(s, sp, idx, total)
        if len(imgs) >= 2:
            return self._content_multi_image(s, sp, items, body, imgs, idx, total)
        if imgs:
            p, a = imgs[0]
            n_words = sum(len((i["head"] + " " + i["text"]).split()) for i in items) + len(body.split())
            if a >= 2.0 and n_words < 70:
                return self._content_band(s, sp, items, body, p, a, idx, total)
            if a < 1.25:
                return self._content_bleed(s, sp, items, body, p, a, idx, total)
            return self._content_framed(s, sp, items, body, p, a, idx, total)
        # text only
        if not items:
            return self._statement(s, sp, body, idx, total)
        heads = sum(1 for i in items if i["head"] and i["text"])
        avg_words = sum(len((i["head"] + " " + i["text"]).split()) for i in items) / len(items)
        want = sp.get("variant")
        if want in ("cards", "grid") or (want is None and heads == len(items) and 2 <= len(items) <= 6
                                          and avg_words <= 34 and self._prev_variant != "cards"):
            self._prev_variant = "cards"
            y = self._header(s, sp)
            if body:
                y = self._lead(s, body, MX, y, SW - 2 * MX)
            self._cards(s, items, MX, y, SW - 2 * MX, SH - MB - y, numbered=len(items) > 2)
        else:
            self._prev_variant = "list"
            if len(items) >= 4 and avg_words <= 12 and not body:
                # two balanced columns
                y = self._header(s, sp)
                self._list_grid(s, items, MX, y, SW - 2 * MX, SH - MB - y, cols=2)
            elif body and len(body.split()) > 25 and len(items) <= 5:
                # lead paragraph left, list right
                y = self._header(s, sp)
                lw = (SW - 2 * MX) * 0.42
                h = SH - MB - y
                ls_ = fit_size(body, self.bf, lw - 0.3, h, self.d["lead_max"] + 2, 13, ls=1.15)
                box(s, MX, y, lw, h, fill=t["surface"], radius=0.14, shadow=not t["dark"])
                text(s, MX + 0.3, y + 0.3, lw - 0.6, h - 0.6, body, font=self.bf, size=min(ls_, self.d["lead_max"]),
                     color=t["text"], ls=1.15, anchor="m")
                self._list(s, items, MX + lw + 0.5, y, SW - 2 * MX - lw - 0.5, h,
                           marker="num" if sp.get("ordered") else "dot")
            else:
                y = self._header(s, sp)
                if body:
                    y = self._lead(s, body, MX, y, SW - 2 * MX)
                self._list(s, items, MX, y, SW - 2 * MX - 0.4, SH - MB - y,
                           marker="num" if sp.get("ordered") else "bar")
        self._footer(s, idx, total)

    def _lead(self, s, body, x, y, w, max_h=1.4):
        ls_ = fit_size(body, self.bf, w, max_h, self.d["lead_max"], 13, ls=1.1)
        bh = para_h(body, self.bf, ls_, w, 1.1)
        text(s, x, y - 0.1, w, bh + 0.05, body, font=self.bf, size=ls_, color=self.t["muted"], ls=1.1)
        return y - 0.1 + bh + 0.35

    def _statement(self, s, sp, body, idx, total):
        t = self.t
        y = self._header(s, sp)
        w = SW - 2 * MX - 1.2
        h = SH - MB - y - 0.2
        size = fit_size(body, self.hf, w, h, 34 if self.d is DENSITY["light"] else 28, 16, ls=1.12)
        bh = para_h(body, self.hf, size, w, 1.12)
        yy = y + max(0.0, (h - bh) / 2 - 0.2)
        box(s, MX, yy + 0.05, 0.08, bh - 0.1, fill=t["accent"])
        text(s, MX + 0.5, yy, w, bh + 0.1, body, font=self.hf, size=size, color=t["text"], ls=1.12)
        self._footer(s, idx, total)

    def _text_column(self, s, sp, items, body, x, w, top=None, bottom=None):
        """Header + lead + list inside a column. Returns nothing."""
        y = self._header(s, sp, x=x, w=w) if top is None else top
        bottom = bottom or SH - MB
        if body and items:
            y = self._lead(s, body, x, y, w, max_h=1.3)
        if items:
            self._list(s, items, x, y, w, bottom - y, marker="num" if sp.get("ordered") else "dot")
        elif body:
            size = fit_size(body, self.bf, w, bottom - y, self.d["lead_max"] + 4, 13, ls=1.15)
            bh = para_h(body, self.bf, size, w, 1.15)
            text(s, x, y + max(0, (bottom - y - bh) / 2 - 0.3), w, bh + 0.1, body, font=self.bf,
                 size=size, color=self.t["text"], ls=1.15)

    def _content_bleed(self, s, sp, items, body, p, a, idx, total):
        """Portrait/square image as a full-height panel flush to one edge."""
        pw = min(max(SH * a, SW * 0.34), SW * 0.46)
        right = self._pick_side(sp)
        if right:
            place_image(s, p, SW - pw, 0, pw, SH, a)
            x, w = MX, SW - pw - MX - 0.6
        else:
            place_image(s, p, 0, 0, pw, SH, a)
            x, w = pw + 0.6, SW - pw - 0.6 - MX
        self._text_column(s, sp, items, body, x, w)
        self._footer_min(s, idx, x, w)

    def _pick_side(self, sp) -> bool:
        """True = image on the right. Honours an explicit image_side, else alternates."""
        want = (sp.get("image_side") or "").lower()
        if want in ("left", "right"):
            return want == "right"
        right = self._side == 1
        self._side *= -1
        return right

    def _footer_min(self, s, idx, x, w):
        text(s, x + w - 1, SH - 0.45, 1, 0.25, f"{idx:02d}", font=self.bf_bold, size=9,
             color=self.t["muted"], align="r")

    def _content_framed(self, s, sp, items, body, p, a, idx, total):
        """Landscape image framed at its natural aspect next to text (zero crop)."""
        t = self.t
        y = self._header(s, sp)
        h = SH - MB - y
        cw = SW - 2 * MX
        n_words = sum(len((i["head"] + " " + i["text"]).split()) for i in items) + len(body.split())
        frac = 0.58 if n_words < 45 else 0.5 if n_words < 90 else 0.44
        iw = min(cw * frac, h * a)
        ih = iw / a
        right = self._pick_side(sp)
        ix = MX + cw - iw if right else MX
        iy = y + (h - ih) / 2
        place_image(s, p, ix, iy, iw, ih, a, radius=0.14, shadow=not t["dark"])
        tx = MX if right else MX + iw + 0.55
        tw = cw - iw - 0.55
        self._text_column(s, sp, items, body, tx, tw, top=y)
        self._footer(s, idx, total)

    def _content_band(self, s, sp, items, body, p, a, idx, total):
        """Panoramic image as a band under the title, text in columns below."""
        t = self.t
        y = self._header(s, sp)
        cw = SW - 2 * MX
        avail = SH - MB - y
        ih = min(cw / a, avail * 0.5)
        iw = ih * a
        place_image(s, p, MX + (cw - iw) / 2, y, iw, ih, a, radius=0.12)
        y2 = y + ih + 0.35
        h2 = SH - MB - y2
        if items and len(items) <= 4 and all(i["head"] for i in items):
            n = len(items)
            colw = (cw - GAP * (n - 1)) / n
            size = min(fit_size(i["text"] or i["head"], self.bf, colw, h2 - 0.45, self.d["body_max"] - 2, 11, ls=1.05)
                       for i in items)
            for k, it in enumerate(items):
                self._card_text(s, it, MX + k * (colw + GAP), y2, colw, size)
        elif items:
            self._list(s, items, MX, y2, cw, h2, valign="t")
        elif body:
            self._lead(s, body, MX, y2 + 0.1, cw, max_h=h2)
        self._footer(s, idx, total)

    def _content_multi_image(self, s, sp, items, body, imgs, idx, total):
        t = self.t
        if not items and not body:
            return self._s_gallery(s, sp, idx, total)
        y = self._header(s, sp)
        h = SH - MB - y
        cw = SW - 2 * MX
        tw = cw * 0.42
        ix, iw = MX + tw + 0.55, cw - tw - 0.55
        for (p, a), (bx, by, bw, bh) in zip(imgs, justified_rows([a for _, a in imgs], iw, h, 0.15)):
            place_image(s, p, ix + bx, y + by, bw, bh, a, radius=0.1)
        self._text_column(s, sp, items, body, MX, tw, top=y)
        self._footer(s, idx, total)

    # ---- gallery / image -----------------------------------------------------
    def _s_gallery(self, s, sp, idx, total):
        imgs = self._images(sp)
        if not imgs:
            return self._s_content(s, sp, idx, total)
        y = self._header(s, sp) if sp.get("title") else MT
        h = SH - MB - y
        w = SW - 2 * MX
        caps = sp.get("captions") or []
        cap_h = 0.35 if caps else 0
        for k, ((p, a), (bx, by, bw, bh)) in enumerate(zip(imgs, justified_rows([a for _, a in imgs], w, h - cap_h, 0.18))):
            place_image(s, p, MX + bx, y + by, bw, bh, a, radius=0.1, shadow=not self.t["dark"])
            if k < len(caps) and caps[k]:
                text(s, MX + bx, y + by + bh + 0.06, bw, 0.3, caps[k], font=self.bf, size=11,
                     color=self.t["muted"], align="c")
        self._footer(s, idx, total)

    def _s_image(self, s, sp, idx, total):
        """Showcase: one hero image, title overlaid, optional caption."""
        imgs = self._images(sp)
        if not imgs:
            return self._s_content(s, sp, idx, total)
        if len(imgs) > 1:
            return self._s_gallery(s, sp, idx, total)
        p, a = imgs[0]
        items = _items(sp)
        if items:
            return self._s_content(s, sp, idx, total)
        cap = sp.get("body") or sp.get("subtitle") or ""
        if a >= 1.3:
            place_image(s, p, 0, 0, SW, SH, a)
            gradient_box(s, 0, SH * 0.4, SW, SH * 0.6, "07070B", 0.0, 0.85, angle=90)
            w = SW - 2 * MX
            ts = fit_size(sp.get("title", ""), self.hf, w, 1.4, 44, 26, max_lines=2)
            th = para_h(sp.get("title", ""), self.hf, ts, w)
            ch = para_h(cap, self.bf, 16, w * 0.8) if cap else 0
            yb = SH - 0.7 - ch - (0.15 if cap else 0)
            text(s, MX, yb - th, w, th + 0.05, sp.get("title", ""), font=self.hf, size=ts, color="FFFFFF")
            if cap:
                text(s, MX, yb + 0.15, w * 0.8, ch + 0.05, cap, font=self.bf, size=16, color="D4D4D8")
        else:
            self._content_bleed(s, {**sp, "body": cap}, [], cap, p, a, idx, total)

    # ---- stats ---------------------------------------------------------------
    def _s_stats(self, s, sp, idx, total):
        t = self.t
        stats = [x for x in (sp.get("stats") or []) if isinstance(x, dict) and x.get("value")]
        if not stats:
            return self._s_content(s, sp, idx, total)
        imgs = self._images(sp)
        y = self._header(s, sp)
        body = sp.get("body") or ""
        x, w = MX, SW - 2 * MX
        if imgs and len(stats) <= 3:
            p, a = imgs[0]
            iw = min(w * 0.42, (SH - MB - y) * a)
            place_image(s, p, MX + w - iw, y, iw, SH - MB - y, a, radius=0.14)
            w = w - iw - 0.5
        if body:
            y = self._lead(s, body, x, y, w, max_h=1.0)
        h = SH - MB - y
        n = len(stats)
        cols = n if n <= 4 else math.ceil(n / 2)
        if w < 8 and n > 2:
            cols = 1 if n <= 3 else 2
        rows = math.ceil(n / cols)
        g = 0.25
        cw = (w - g * (cols - 1)) / cols
        ch_max = (h - g * (rows - 1)) / rows
        pad = 0.3
        iw_ = cw - 2 * pad
        longest = max((str(st.get("value")) for st in stats), key=lambda v: text_w(v, self.hf, 10))
        vs = fit_size(longest, self.hf, iw_, ch_max * 0.45, 60 if rows == 1 else 44, 22, max_lines=1)
        lab_size = min(self.d["body_max"] - 2, 20)
        desc_size = self.d["body_max"] - 4
        # tile height = tallest content, so tiles never look half-empty
        need = 0.0
        for st in stats:
            nh = line_h(self.hf, vs)
            if st.get("label"):
                nh += para_h(str(st["label"]), self.bf_bold, lab_size, iw_) + 0.06
            if st.get("desc"):
                nh += para_h(str(st["desc"]), self.bf, desc_size, iw_, 1.05)
            need = max(need, nh)
        ch = min(ch_max, max(need + 2 * pad + 0.1, min(ch_max, 2.0)))
        top = _vpos(y, h, rows * ch + (rows - 1) * g)
        for i, st in enumerate(stats):
            r, c = divmod(i, cols)
            in_row = min(cols, n - r * cols)
            off = (w - (in_row * cw + (in_row - 1) * g)) / 2
            cx, cy = x + off + c * (cw + g), top + r * (ch + g)
            acc = t["series"][i % 3]
            box(s, cx, cy, cw, ch, fill=t["surface"], radius=0.14, shadow=not t["dark"],
                line=t["line"] if t["dark"] else None)
            box(s, cx, cy + 0.25, 0.07, ch - 0.5, fill=acc)
            vh = line_h(self.hf, vs)
            text(s, cx + pad, cy + pad - 0.05, cw - 2 * pad, vh, str(st["value"]), font=self.hf, size=vs,
                 color=acc if _contrast(acc, t["surface"]) > 2.5 else t["accent_text"])
            yy = cy + pad + vh
            lab = str(st.get("label") or "")
            if lab:
                ls_ = fit_size(lab, self.bf_bold, cw - 2 * pad, 0.8, lab_size, 11, max_lines=2)
                lh = para_h(lab, self.bf_bold, ls_, cw - 2 * pad)
                text(s, cx + pad, yy, cw - 2 * pad, lh + 0.02, lab, font=self.bf_bold, size=ls_, color=t["text"])
                yy += lh + 0.06
            desc = str(st.get("desc") or "")
            if desc:
                room = cy + ch - pad - yy
                ds = fit_size(desc, self.bf, cw - 2 * pad, room, desc_size, 10, ls=1.05)
                text(s, cx + pad, yy, cw - 2 * pad, room, desc, font=self.bf, size=ds, color=t["muted"], ls=1.05)
        self._footer(s, idx, total)

    # ---- cards ---------------------------------------------------------------
    def _s_cards(self, s, sp, idx, total):
        self._prev_variant = "cards"
        items = _items(sp)
        if not items:
            return self._s_content(s, sp, idx, total)
        if self._images(sp):
            return self._s_content(s, sp, idx, total)
        y = self._header(s, sp)
        body = sp.get("body") or ""
        if body:
            y = self._lead(s, body, MX, y, SW - 2 * MX)
        self._cards(s, items, MX, y, SW - 2 * MX, SH - MB - y, numbered=sp.get("numbered", True))
        self._footer(s, idx, total)

    # ---- process -------------------------------------------------------------
    def _s_process(self, s, sp, idx, total):
        t = self.t
        items = _items(sp)
        if not items:
            return self._s_content(s, sp, idx, total)
        n = len(items)
        max_words = max(len((i["head"] + " " + i["text"]).split()) for i in items)
        y = self._header(s, sp)
        body = sp.get("body") or ""
        if body:
            y = self._lead(s, body, MX, y, SW - 2 * MX)
        w = SW - 2 * MX
        h = SH - MB - y
        if n <= 5 and max_words <= 30 and not self._images(sp):
            g = 0.3
            cw = (w - g * (n - 1)) / n
            d = 0.7
            size = min(fit_size(i["text"] or i["head"], self.bf, cw - 0.1, h - d - 1.0,
                                self.d["body_max"] - 1, 11, ls=1.05) for i in items)
            txt_h = max((para_h(i["head"], self.bf_bold, size + 3, cw - 0.1, bold_family=self.bf_bold) + 0.08
                         if i["head"] else 0) + (para_h(i["text"], self.bf, size, cw - 0.1, 1.05) if i["text"] else 0)
                        for i in items)
            cy = _vpos(y, h, d + 0.3 + txt_h)
            line_shape(s, MX + cw / 2, cy + d / 2, MX + w - cw / 2, cy + d / 2, t["line"], 2.0)
            txt_top = cy + d + 0.3
            for k, it in enumerate(items):
                cx = MX + k * (cw + g)
                acc = t["series"][k % 3]
                box(s, cx + cw / 2 - d / 2, cy, d, d, fill=acc, shape=MSO_SHAPE.OVAL)
                text(s, cx + cw / 2 - d / 2, cy, d, d, str(k + 1), font=self.hf, size=22,
                     color=on_color(acc), align="c", anchor="m")
                if k < n - 1:
                    ax = cx + cw + g / 2
                    box(s, ax - 0.09, cy + d / 2 - 0.09, 0.18, 0.18, fill=t["muted"], shape=MSO_SHAPE.CHEVRON)
                yy = txt_top
                if it["head"]:
                    hs = size + 3
                    hh = para_h(it["head"], self.bf_bold, hs, cw - 0.1, bold_family=self.bf_bold)
                    text(s, cx + 0.05, yy, cw - 0.1, hh + 0.02, it["head"], font=self.bf_bold, size=hs,
                         color=t["text"], align="c")
                    yy += hh + 0.08
                if it["text"]:
                    text(s, cx + 0.05, yy, cw - 0.1, SH - MB - yy, it["text"], font=self.bf, size=size,
                         color=t["muted"], align="c", ls=1.05)
        else:
            imgs = self._images(sp)
            lw = w
            if imgs:
                p, a = imgs[0]
                iw = min(w * 0.4, h * a)
                place_image(s, p, MX + w - iw, y + (h - iw / a) / 2 if iw / a < h else y, iw, min(h, iw / a), a, radius=0.14)
                lw = w - iw - 0.5
            cols = 2 if n >= 6 and not imgs else 1
            per = math.ceil(n / cols)
            cw = (lw - 0.5 * (cols - 1)) / cols
            for c in range(cols):
                self._list(s, items[c * per:(c + 1) * per], MX + c * (cw + 0.5), y, cw, h, marker="num", valign="t")
        self._footer(s, idx, total)

    # ---- timeline ------------------------------------------------------------
    def _s_timeline(self, s, sp, idx, total):
        t = self.t
        items = _items(sp)
        if not items:
            return self._s_content(s, sp, idx, total)
        n = len(items)
        if n > 7:
            return self._s_process(s, {**sp, "kind": "process"}, idx, total)
        y = self._header(s, sp)
        w = SW - 2 * MX
        h = SH - MB - y
        mid = y + h * 0.42 if n >= 5 else _vpos(y, h, 2.2) + 0.3
        line_shape(s, MX, mid, MX + w, mid, t["line"], 2.5)
        slot = w / n
        alt = n >= 5
        for k, it in enumerate(items):
            cx = MX + slot * k + slot / 2
            acc = t["series"][k % 3]
            box(s, cx - 0.14, mid - 0.14, 0.28, 0.28, fill=acc, shape=MSO_SHAPE.OVAL,
                line=t["bg"], line_w=3)
            label = it.get("date") or it["head"]
            detail_head = it["head"] if it.get("date") else ""
            detail = it["text"]
            tw = slot * (1.8 if alt else 0.94)
            tw = min(tw, 2 * (cx - MX) + slot * 0.1, 2 * (MX + w - cx) + slot * 0.1)   # stay centred on the dot
            tx = cx - tw / 2
            up = alt and k % 2 == 0
            if up:
                room = mid - y - 0.35
                paras, hh = self._tl_block(label, detail_head, detail, tw, room, acc)
                text(s, tx, mid - 0.3 - hh, tw, hh + 0.05, paras, font=self.bf, size=12, color=t["muted"],
                     align="c", anchor="b")
            else:
                room = SH - MB - mid - 0.35
                paras, hh = self._tl_block(label, detail_head, detail, tw, room, acc)
                text(s, tx, mid + 0.3, tw, hh + 0.05, paras, font=self.bf, size=12, color=t["muted"], align="c")
        self._footer(s, idx, total)

    def _tl_block(self, label, head, detail, w, room, acc):
        t = self.t
        for size in _sizes(self.d["body_max"] - 2, 10):
            ls_ = size + 4
            h = line_h(self.hf, ls_) * count_lines(label, self.hf, ls_, w)
            if head:
                h += para_h(head, self.bf_bold, size + 1, w, bold_family=self.bf_bold)
            if detail:
                h += para_h(detail, self.bf, size, w, 1.05)
            if h <= room:
                break
        paras = [[(label, {"font": self.hf, "size": ls_, "color": acc if _contrast(acc, t["bg"]) > 2.5 else t["accent_text"]})]]
        if head:
            paras.append([(head, {"font": self.bf_bold, "size": size + 1, "color": t["text"]})])
        if detail:
            paras.append([(detail, {"size": size})])
        return paras, h

    # ---- comparison ----------------------------------------------------------
    def _s_comparison(self, s, sp, idx, total):
        t = self.t
        cols = [c for c in (sp.get("columns") or []) if isinstance(c, dict)]
        if len(cols) < 2:
            return self._s_content(s, sp, idx, total)
        cols = cols[:4]
        y = self._header(s, sp)
        w = SW - 2 * MX
        h = SH - MB - y
        n = len(cols)
        g = 0.7 if n == 2 else 0.3
        cw = (w - g * (n - 1)) / n
        hdr_h = 0.7
        pad = 0.3
        tw = cw - 2 * pad - 0.35
        pts_all = [[str(p) for p in (c.get("points") or [])] for c in cols]
        size, need = self.d["body_min"], h
        for sz in _sizes(self.d["body_max"], max(10, self.d["body_min"] - 2)):
            need = max((sum(para_h(p, self.bf, sz, tw, 1.05) for p in pts) + 0.18 * max(len(pts) - 1, 0))
                       for pts in pts_all)
            if need <= h - hdr_h - 2 * pad:
                size = sz
                break
        ch = min(h, max(hdr_h + 2 * pad + need + 0.1, min(h, 2.8)))
        y = _vpos(y, h, ch)
        h = ch
        for k, (c, pts) in enumerate(zip(cols, pts_all)):
            cx = MX + k * (cw + g)
            acc = t["series"][k % 3] if n > 2 else (t["muted"] if k == 0 and sp.get("contrast") else t["series"][k])
            box(s, cx, y, cw, h, fill=t["surface"], radius=0.14, shadow=not t["dark"],
                line=t["line"] if t["dark"] else None)
            box(s, cx, y, cw, hdr_h, fill=acc, radius=0.14)
            box(s, cx, y + hdr_h - 0.16, cw, 0.16, fill=acc)
            text(s, cx + pad, y, cw - 2 * pad, hdr_h, c.get("heading", ""), font=self.bf_bold,
                 size=fit_size(c.get("heading", ""), self.bf_bold, cw - 2 * pad, hdr_h - 0.1, 20, 12, max_lines=2),
                 color=on_color(acc), anchor="m")
            yy = y + hdr_h + pad
            for ptxt in pts:
                ph = para_h(ptxt, self.bf, size, tw, 1.05)
                dd = 0.11
                box(s, cx + pad + 0.05, yy + line_h(self.bf, size) / 2 - dd / 2, dd, dd, fill=acc, shape=MSO_SHAPE.OVAL)
                text(s, cx + pad + 0.35, yy, tw, ph + 0.02, ptxt, font=self.bf, size=size, color=t["text"], ls=1.05)
                yy += ph + 0.18
        if n == 2:
            vx, vy = MX + cw + g / 2 - 0.32, y + h / 2 - 0.32
            box(s, vx, vy, 0.64, 0.64, fill=t["bg"], shape=MSO_SHAPE.OVAL, line=t["line"], line_w=1.5)
            text(s, vx, vy, 0.64, 0.64, "VS", font=self.bf_bold, size=13, color=t["text"], align="c", anchor="m")
        self._footer(s, idx, total)

    # ---- table ---------------------------------------------------------------
    def _s_table(self, s, sp, idx, total):
        t = self.t
        tb = sp.get("table") or {}
        header = [str(x) for x in (tb.get("header") or [])]
        rows = [[str(c) for c in r] for r in (tb.get("rows") or []) if isinstance(r, (list, tuple))]
        if not header and rows:
            header, rows = rows[0], rows[1:]
        if not header:
            return self._s_content(s, sp, idx, total)
        ncol = len(header)
        rows = [(r + [""] * ncol)[:ncol] for r in rows][:14]
        y = self._header(s, sp)
        w = SW - 2 * MX
        h = SH - MB - y
        # column widths ∝ content length (clamped)
        lens = [max([len(header[c])] + [len(r[c]) for r in rows]) for c in range(ncol)]
        lens = [min(max(L, 6), 40) for L in lens]
        cws = [w * L / sum(lens) for L in lens]
        size, rh, hh = 10, [0.4] * len(rows), 0.45
        top_size = 22 if self.d is DENSITY["light"] or self.deck.get("density") == "light" else 18
        for sz in _sizes(top_size, 9):
            rh = [max(para_h(r[c], self.bf, sz, cws[c] - 0.25) for c in range(ncol)) + 0.2 for r in rows]
            hh = max(para_h(hd, self.bf_bold, sz, cws[c] - 0.25) for c, hd in enumerate(header)) + 0.24
            if sum(rh) + hh <= h:
                size = sz
                break
        # spread rows a little when there is room, but keep the table compact
        spare = h - sum(rh) - hh
        extra = max(0.0, min(spare / (len(rows) + 1), 0.25))
        shape = s.shapes.add_table(len(rows) + 1, ncol, E(MX), E(y), E(w), E(sum(rh) + hh))
        tbl = shape.table
        tbl.rows[0].height = E(hh + extra)
        for r in range(len(rows)):
            tbl.rows[r + 1].height = E(rh[r] + extra)
        # strip default table style banding look
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "1")
        for c in range(ncol):
            tbl.columns[c].width = E(cws[c])
        for r in range(len(rows) + 1):
            for c in range(ncol):
                cell = tbl.cell(r, c)
                val = header[c] if r == 0 else rows[r - 1][c]
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(t["accent"])
                else:
                    cell.fill.fore_color.rgb = _rgb(t["surface"] if r % 2 else t["bg"] if not t["dark"] else t["surface2"])
                cell.margin_left = cell.margin_right = E(0.1)
                cell.margin_top = cell.margin_bottom = E(0.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = ""
                run = tf.paragraphs[0].add_run()
                run.text = val
                run.font.size = Pt(size)
                run.font.name = self.bf_bold if r == 0 or c == 0 else self.bf
                run.font.color.rgb = _rgb(t["on_accent"] if r == 0 else t["text"])
        self._footer(s, idx, total)

    # ---- chart ---------------------------------------------------------------
    def _s_chart(self, s, sp, idx, total):
        t = self.t
        ch = sp.get("chart") or {}
        labels = [str(x) for x in (ch.get("labels") or [])]
        series = ch.get("series")
        if not series and ch.get("values"):
            series = [{"name": ch.get("name") or sp.get("title", "Series"), "values": ch.get("values")}]
        clean = []
        for se in series or []:
            try:
                vals = [float(re.sub(r"[^\d.\-]", "", str(v)) or 0) for v in se.get("values", [])][:len(labels)]
                if vals:
                    clean.append((str(se.get("name", "")), vals))
            except Exception:
                continue
        if not labels or not clean:
            return self._s_content(s, sp, idx, total)
        y = self._header(s, sp)
        items = _items(sp)
        w = SW - 2 * MX
        h = SH - MB - y
        cw = w
        if items:
            cw = w * 0.6
            self._list(s, items, MX + cw + 0.5, y, w - cw - 0.5, h, marker="bar")
        box(s, MX, y, cw, h, fill=t["surface"], radius=0.14, shadow=not t["dark"], line=t["line"] if t["dark"] else None)
        ctype = str(ch.get("type", "column")).lower()
        xl = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "horizontal_bar": XL_CHART_TYPE.BAR_CLUSTERED,
              "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
              "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "donut": XL_CHART_TYPE.DOUGHNUT,
              "area": XL_CHART_TYPE.AREA, "stacked": XL_CHART_TYPE.COLUMN_STACKED}.get(ctype, XL_CHART_TYPE.COLUMN_CLUSTERED)
        if xl in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
            clean = clean[:1]
        cd = CategoryChartData()
        cd.categories = labels
        for name, vals in clean:
            cd.add_series(name, vals + [0] * (len(labels) - len(vals)))
        gf = s.shapes.add_chart(xl, E(MX + 0.2), E(y + 0.2), E(cw - 0.4), E(h - 0.4), cd)
        chart = gf.chart
        chart.font.size = Pt(12)
        chart.font.name = self.bf
        chart.font.color.rgb = _rgb(t["muted"])
        pie = xl in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
        chart.has_title = False
        chart.has_legend = pie or len(clean) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.RIGHT if pie else XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.color.rgb = _rgb(t["text"])
        plot = chart.plots[0]
        unit = str(ch.get("unit") or "")
        try:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(11)
            dl.font.color.rgb = _rgb(t["text"] if not pie else "FFFFFF")
            ints = all(float(v).is_integer() for _, vals in clean for v in vals)
            base = "#,##0" if ints else "#,##0.0"
            dl.number_format = base + (f'"{unit}"' if unit and len(unit) <= 3 else "")
            dl.number_format_is_linked = False
            if not pie and xl != XL_CHART_TYPE.LINE_MARKERS:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception:
            pass
        if pie:
            pts = plot.series[0].points
            for i in range(len(labels)):
                pts[i].format.fill.solid()
                pts[i].format.fill.fore_color.rgb = _rgb(t["series"][i % len(t["series"])])
        else:
            try:
                plot.gap_width = 70
            except Exception:
                pass
            for i, se in enumerate(plot.series):
                col = t["series"][i % len(t["series"])]
                if xl == XL_CHART_TYPE.LINE_MARKERS:
                    se.format.line.color.rgb = _rgb(col)
                    se.format.line.width = Pt(3)
                    se.smooth = False
                else:
                    se.format.fill.solid()
                    se.format.fill.fore_color.rgb = _rgb(col)
            for ax in (chart.category_axis, chart.value_axis):
                ax.format.line.color.rgb = _rgb(t["line"])
                ax.tick_labels.font.color.rgb = _rgb(t["muted"])
                ax.tick_labels.font.size = Pt(11)
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(t["line"])
            chart.value_axis.visible = False
        self._footer(s, idx, total)


def _vpos(y: float, h: float, block: float) -> float:
    """Optical centre: put a block slightly above the middle of the free area under the title."""
    return y + max(0.0, (h - block) * 0.42)


def _as_plain_content(sp: dict) -> dict:
    """Last-resort fallback: every word of a composite slide as plain bullets (never an empty divider)."""
    sp = _deep_plain(sp)
    items = list(sp.get("bullets") or [])
    for sec in sp.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        for it in sec.get("items") or []:
            if isinstance(it, dict):
                head = it.get("head") or it.get("value") or ""
                txt = it.get("text") or " ".join(x for x in (it.get("label"), it.get("desc")) if x)
                if head or txt:
                    items.append({"head": head, "text": txt})
        for row in sec.get("rows") or []:
            if isinstance(row, (list, tuple)) and row:
                items.append({"head": str(row[0]), "text": " → ".join(str(c) for c in row[1:])})
    return {**sp, "kind": "content", "bullets": items, "body": sp.get("body") or sp.get("lead") or ""}


def _deep_plain(o):
    if isinstance(o, str):
        return plain(o)
    if isinstance(o, list):
        return [_deep_plain(x) for x in o]
    if isinstance(o, dict):
        return {k: (_deep_plain(v) if k not in ("images", "theme") else v) for k, v in o.items()}
    return o


def _sizes(hi, lo):
    s = float(hi)
    while s >= lo:
        yield s
        s -= 1 if s > 14 else 0.5
