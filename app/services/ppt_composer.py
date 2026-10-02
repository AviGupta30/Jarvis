"""
ppt_composer.py — Composite "infographic" slides for PPT v6 (hackathon / pitch grade)
======================================================================================
Renders a slide spec of kind "sections" the way a senior designer builds a dense
pitch slide (reference: SIH idea decks):

    title + optional subtitle
    lead statement (bold, key words in accent)            │  image panel
    ┌ SECTION PILL ┐                                      │  (pill heading, framed
    icon cards / list / steps / stats / fields / table     │   screenshots, numbered
    ┌ SECTION PILL ┐ …                                    │   captions)
    ────────────────────────────────────────────────────────────────────────
    ┌ HIGHLIGHT PILL ┐  tagline        (bottom band: "Why we stand out" etc.)
    tinted card · tinted card · tinted card · tinted card
    ═══════════════════ footer bar ═══════════════════ page ═

Everything is measured: one body font size is chosen for the whole slide (the
largest that fits every zone), zones get the height they need, and leftover
space is distributed so nothing floats or overflows. **bold** markup in the
text is rendered as bold runs.

Spec:
  {"kind": "sections", "title", "subtitle", "lead",
   "sections": [{"heading", "style": cards|list|steps|stats|fields|table|paragraph|gallery,
                 "highlight": bool, "lead", "items": [{"head", "text", "value", "label"}],
                 "header": [...], "rows": [[...]]}],
   "images": [...]}
"""
from __future__ import annotations

import os
import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

from app.services.ppt_designer import (SW, SH, MX, MT, E, box, text, place_image, justified_rows, fit_size,
                                       para_h, line_h, count_lines, text_w, rich, plain, icon_for, icon_png,
                                       on_color, _mix, _contrast, _soft_shadow, _sizes, _rgb)

FOOT_H = 0.3
_COMPACT = [False]          # set by the planner while measuring/drawing one slide
PILL_H = 0.32
GAP = 0.24
CMX = 0.42          # tight side margin for composite slides (reference decks run nearly edge to edge)
SEC_GAP = 0.16
LS = 1.0

_HIGHLIGHT = re.compile(r"why|stand out|advantage|benefit|usp|highlight|takeaway|differentiat|unique|key point|"
                        r"impact|outcome|value", re.I)
_GALLERY = re.compile(r"prototype|screenshot|demo|preview|gallery|mock|ui|interface|screens|product tour|snapshots",
                      re.I)


# ══════════════════════════════════════════════════════════════════════════════
#  helpers
# ══════════════════════════════════════════════════════════════════════════════
# Pastel card family (as in top hackathon decks): the theme accent leads, harmonised pastels follow.
_PASTEL_INK = ["B7791F", "B83268", "2E8B57", "6D4AB8", "1B7F95"]


def _inks(t: dict) -> list[str]:
    if t["dark"]:
        return [t["accent"], t["accent2"], t["accent3"], "F59E0B", "F472B6", "34D399"]
    return [t["accent"]] + _PASTEL_INK


def _tints(t: dict) -> list[str]:
    if t["dark"]:
        return [_mix(c, t["bg"], 0.8) for c in _inks(t)]
    return [_mix(c, "FFFFFF", 0.88) for c in _inks(t)]


def _ink(t: dict, k: int) -> str:
    """Strong colour for icons / numbers on the k-th tint."""
    c = _inks(t)[k % 6]
    bg = _tints(t)[k % 6]
    for _ in range(5):
        if _contrast(c, bg) >= 3.2:
            break
        c = _mix(c, t["text"], 0.3)
    return c


def _pill_color(t: dict) -> str:
    c = _mix(t["accent"], t["accent2"], 0.35)
    return c if _contrast(c, "FFFFFF") >= 3 or t["dark"] else _mix(c, "000000", 0.25)


def _item_runs(it: dict) -> list[tuple[str, bool]]:
    head, txt = plain(it.get("head") or ""), it.get("text") or ""
    if head and txt:
        return [(head + ": ", True)] + rich(txt)
    return rich(head or txt) if not head else [(head, True)]


def _to_paras(runs, font, bold_font, color, bold_color=None, size=None):
    return [[(seg, {"bold": b, "font": bold_font if b else font, "color": (bold_color or color) if b else color,
                    **({"size": size} if size else {})}) for seg, b in runs]]


def _h(r, runs, size, w, ls=LS):
    return count_lines(runs, r.bf, size, w, bold_family=r.bf_bold) * line_h(r.bf, size, ls)


def _lead_h(r, lead: str, size: float, w: float) -> float:
    """Leads are drawn entirely in the semibold heading font — measure them that way."""
    runs = [(seg, True) for seg, _ in rich(lead)]
    return count_lines(runs, r.hf, size, w, bold_family=r.hf) * line_h(r.hf, size, 1.1)


def _pill(r, s, x, y, label, draw=True, color=None):
    """Rounded section heading chip. Returns its width."""
    lab = plain(label).upper()
    size = 12
    w = text_w(lab, r.bf_bold, size, True) * 1.12 + 0.44
    if draw:
        c = color or _pill_color(r.t)
        box(s, x, y, w, PILL_H, fill=c, radius=PILL_H / 2)
        text(s, x, y, w, PILL_H, lab, font=r.bf_bold, size=size, color=on_color(c), align="c", anchor="m", spacing=1)
    return w


def _norm_sections(sp: dict) -> list[dict]:
    out = []
    for sec in sp.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        raw_items = sec.get("items") or []
        if isinstance(raw_items, list) and raw_items and all(
                isinstance(i, str) and i.strip().lower() in ("header", "rows", "table", "columns", "data")
                for i in raw_items):
            raw_items = []                                # a table whose data was lost upstream → no fake rows
        tbl = sec.get("table") if isinstance(sec.get("table"), dict) else (raw_items if isinstance(raw_items, dict) else None)
        if tbl:                                   # {"items": {"header": […], "rows": […]}} / {"table": {…}}
            sec = {**sec, "style": "table", "header": tbl.get("header") or sec.get("header") or [],
                   "rows": tbl.get("rows") or sec.get("rows") or []}
            raw_items = [] if isinstance(raw_items, dict) else raw_items
        # an item that is really a heading ("Before vs. After X:") starts a new section of its own
        split_at = [k for k, it in enumerate(raw_items) if isinstance(it, dict) and not str(it.get("text") or "").strip()
                    and str(it.get("head") or "").strip().endswith(":") and len(str(it.get("head")).split()) <= 8] + \
                   [k for k, it in enumerate(raw_items) if isinstance(it, dict) and not str(it.get("head") or "").strip()
                    and str(it.get("text") or "").strip().endswith(":") and len(str(it.get("text")).split()) <= 8]
        if split_at and isinstance(raw_items, list):
            k = min(split_at)
            head_it = raw_items[k]
            new_head = (str(head_it.get("head") or "") or str(head_it.get("text") or "")).strip().rstrip(":")
            rest = {**sec, "heading": new_head, "items": raw_items[k + 1:], "highlight": False, "lead": ""}
            sp_tail = {"sections": [rest]}
            raw_items = raw_items[:k]
            tail = _norm_sections(sp_tail)
        else:
            tail = []
        items = []
        for it in raw_items:
            if isinstance(it, str):
                it = {"head": "", "text": it}
            if isinstance(it, dict) and any(str(it.get(k) or "").strip() for k in ("head", "text", "value", "label")):
                items.append({k: str(it.get(k) or "").strip() for k in ("head", "text", "value", "label", "desc")})
        style = str(sec.get("style") or "cards").lower()
        if style not in ("cards", "list", "steps", "stats", "fields", "table", "paragraph", "gallery", "compare"):
            style = "cards"
        if style == "compare":
            style = "table"
        s2 = {"heading": str(sec.get("heading") or "").strip(), "style": style, "items": items,
              "highlight": bool(sec.get("highlight")), "lead": str(sec.get("lead") or "").strip(),
              "header": [str(x) for x in sec.get("header") or []],
              "rows": [[str(c) for c in row] for row in sec.get("rows") or [] if isinstance(row, (list, tuple))]}
        if style == "table" and not s2["rows"] and items:
            s2["header"] = s2["header"] or ["", ""]
            s2["rows"] = [[plain(i["head"]), i["text"]] for i in items]
        if style == "stats" and not all(i.get("value") for i in items):
            for i in items:
                i["value"] = i.get("value") or i.get("head") or ""
                i["label"] = i.get("label") or (i.get("text") if i.get("head") else "")
        if items or s2["rows"] or s2["lead"]:
            prev = out[-1] if out else None
            # a lone card under its own pill wastes a whole row → fold it into the previous card group
            if prev and len(items) == 1 and s2["style"] == "cards" == prev["style"] and not s2["highlight"] \
                    and not prev["highlight"] and not s2["lead"]:
                it = dict(items[0])
                if not it.get("head"):
                    it["head"] = s2["heading"]
                prev["items"].append(it)
            else:
                out.append(s2)
        out += tail
    return out


# ══════════════════════════════════════════════════════════════════════════════
#  section bodies  (each: measure when draw=False, draw when True; returns height)
# ══════════════════════════════════════════════════════════════════════════════
def _cols_for(n: int, w: float) -> int:
    if n <= 1 or w < 4.2:
        return 1
    if w >= 9.5 and n >= 5 or (w >= 9.5 and n == 3):
        return 3
    return 2


def _balanced_rows(items: list, cols: int) -> list[list]:
    """Split items into rows with at most one item difference (5 in 3 cols → 3+2, 7 → 3+2+2).
    Every row is then stretched to the full width, so a grid never has an empty slot."""
    n = len(items)
    if n == 0:
        return []
    cols = max(1, min(cols, n))
    nrows = -(-n // cols)
    base, rem = divmod(n, nrows)
    out, i = [], 0
    for k in range(nrows):
        c = base + (1 if k < rem else 0)
        out.append(items[i:i + c])
        i += c
    return out


def _body_cards(r, s, sec, x, y, w, size, draw, k0=0, cols=None, equal=False, compact=None, extra=0.0):
    t = r.t
    items = sec["items"]
    n = len(items)
    cols = cols or _cols_for(n, w)
    compact = _COMPACT[0] if compact is None else compact
    g = 0.1 if compact else 0.14
    pad = 0.07 if compact else 0.12
    ic = 0.24 if compact else max(0.28, min(0.42, size / 72 * 2.2))
    rows = _balanced_rows(items, cols)
    geo = []                                          # per row: (card width, text width)
    for row in rows:
        cw = (w - g * (len(row) - 1)) / len(row)
        geo.append((cw, cw - 2 * pad - ic - 0.12))
    lh1 = line_h(r.bf, size, LS)
    off = max(0.0, (ic - lh1) / 2)                    # first text line centred on the icon
    heights = [max(max(off + _h(r, _item_runs(it), size, tw), ic) + 2 * pad for it in row)
               for row, (cw, tw) in zip(rows, geo)]
    if equal and heights:
        heights = [max(heights)] * len(heights)
    total = sum(heights) + g * (len(rows) - 1)
    if draw:
        heights = [h + extra / len(rows) for h in heights]      # fill the space the planner gave us
        tints = _tints(t)
        yy = y
        k = k0
        for row, rh, (cw, tw) in zip(rows, heights, geo):
            for c, it in enumerate(row):
                cx = x + c * (cw + g)
                box(s, cx, yy, cw, rh, fill=tints[k % len(tints)], radius=0.1)
                runs = _item_runs(it)
                th = _h(r, runs, size, tw)
                blk = max(ic, off + th)
                top = yy + (rh - blk) / 2                        # icon + text centred as one block
                p = icon_png(icon_for(plain(it.get("head", "") + " " + it.get("text", "")), k), _ink(t, k))
                if p:
                    place_image(s, p, cx + pad, top, ic, ic, 1.0, mode="contain")
                text(s, cx + pad + ic + 0.12, top + off - 0.02, tw, th + 0.05,
                     _to_paras(runs, r.bf, r.bf_bold, t["text"]), font=r.bf, size=size, color=t["text"], ls=LS)
                k += 1
            yy += rh + g
    return total


def _body_list(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    mk = 0.26
    gap = size / 72 * 0.55
    hs = [_h(r, _item_runs(it), size, w - mk) for it in sec["items"]]
    total = sum(hs) + gap * (len(hs) - 1)
    if draw:
        gap += min(0.3, extra / max(1, len(hs) - 1)) if len(hs) > 1 else 0
        yy = y + (extra / 2 if len(hs) == 1 else 0)
        for k, (it, hh) in enumerate(zip(sec["items"], hs)):
            dd = max(0.08, size / 72 * 0.42)
            box(s, x + 0.03, yy + line_h(r.bf, size, 1.05) / 2 - dd / 2, dd, dd, fill=_ink(t, k), shape=MSO_SHAPE.OVAL)
            text(s, x + mk, yy, w - mk, hh + 0.04, _to_paras(_item_runs(it), r.bf, r.bf_bold, t["text"]),
                 font=r.bf, size=size, color=t["text"], ls=LS)
            yy += hh + gap
    return total


def _body_steps(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    items = sec["items"]
    n = len(items)
    if n > 5 or w < 5.5 or (w - 0.22 * (n - 1)) / max(n, 1) < 1.6:      # cards would be too narrow
        # vertical numbered list
        mk = 0.42
        gap = size / 72 * 0.6
        hs = [max(_h(r, _item_runs(it), size, w - mk), 0.3) for it in items]
        total = sum(hs) + gap * (n - 1)
        if draw:
            gap += min(0.3, extra / max(1, n - 1))
            yy = y
            for k, (it, hh) in enumerate(zip(items, hs)):
                d = 0.3
                c = _ink(r.t, k)
                box(s, x, yy, d, d, fill=c, shape=MSO_SHAPE.OVAL)
                text(s, x, yy, d, d, str(k + 1), font=r.bf_bold, size=11, color=on_color(c), align="c", anchor="m")
                text(s, x + mk, yy, w - mk, hh + 0.04, _to_paras(_item_runs(it), r.bf, r.bf_bold, t["text"]),
                     font=r.bf, size=size, color=t["text"], ls=LS)
                yy += hh + gap
        return total
    g = 0.22
    cw = (w - g * (n - 1)) / n
    pad = 0.12
    d = 0.34
    hs = []
    for it in items:
        h = d + 0.1
        if it.get("head"):
            h += _h(r, [(plain(it["head"]), True)], size + 1, cw - 2 * pad) + 0.04
        if it.get("text"):
            h += _h(r, rich(it["text"]), size - 0.5, cw - 2 * pad)
        hs.append(h + 2 * pad)
    total = max(hs)
    if draw:
        tints = _tints(t)
        full = total + extra
        for k, it in enumerate(items):
            cx = x + k * (cw + g)
            box(s, cx, y, cw, full, fill=tints[k % len(tints)], radius=0.1)
            top = y + (full - hs[k]) / 2                 # content block centred in the stretched card
            c = _ink(t, k)
            box(s, cx + pad, top + pad, d, d, fill=c, shape=MSO_SHAPE.OVAL)
            text(s, cx + pad, top + pad, d, d, str(k + 1), font=r.bf_bold, size=12, color=on_color(c), align="c", anchor="m")
            if k < n - 1:
                box(s, cx + cw + g / 2 - 0.08, y + full / 2 - 0.08, 0.16, 0.16, fill=t["muted"], shape=MSO_SHAPE.CHEVRON)
            yy = top + pad + d + 0.1
            if it.get("head"):
                hh = _h(r, [(plain(it["head"]), True)], size + 1, cw - 2 * pad)
                text(s, cx + pad, yy, cw - 2 * pad, hh + 0.03, plain(it["head"]), font=r.bf_bold, size=size + 1, color=t["text"])
                yy += hh + 0.04
            if it.get("text"):
                th = _h(r, rich(it["text"]), size - 0.5, cw - 2 * pad)
                text(s, cx + pad, yy, cw - 2 * pad, th + 0.04, _to_paras(rich(it["text"]), r.bf, r.bf_bold, t["muted"], t["text"]),
                     font=r.bf, size=size - 0.5, color=t["muted"], ls=LS)
    return total


def _body_stats(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    items = sec["items"]
    n = len(items)
    cols = n if n <= 4 else (n + 1) // 2
    g = 0.16
    cw = (w - g * (cols - 1)) / cols
    pad = 0.14
    vals = [plain(i.get("value") or i.get("head") or "") for i in items]
    longest = max(vals, key=lambda v: text_w(v, r.hf, 10)) if vals else ""
    vs = fit_size(longest, r.hf, cw - 2 * pad, 0.8, min(34, size * 2.4), size + 4, max_lines=1)
    hs = []
    for it in items:
        h = line_h(r.hf, vs) + 0.02
        lab = plain(it.get("label") or (it.get("text") if it.get("head") else ""))
        if lab:
            h += _h(r, [(lab, True)], size, cw - 2 * pad)
        if it.get("desc"):
            h += _h(r, rich(it["desc"]), size - 1, cw - 2 * pad)
        hs.append(h + 2 * pad)
    rows = (n + cols - 1) // cols
    rh = max(hs) if hs else 0
    total = rows * rh + g * (rows - 1)
    if draw:
        tints = _tints(t)
        rows_l = _balanced_rows(list(range(n)), cols)
        rh2 = rh + extra / max(1, rows)
        for rr, row in enumerate(rows_l):
            cw_r = (w - g * (len(row) - 1)) / len(row)
            for c, k in enumerate(row):
                it = items[k]
                cx, cy = x + c * (cw_r + g), y + rr * (rh2 + g)
                box(s, cx, cy, cw_r, rh2, fill=tints[k % len(tints)], radius=0.1)
                vh = line_h(r.hf, vs)
                top = cy + (rh2 - hs[k]) / 2
                text(s, cx + pad, top + pad, cw_r - 2 * pad, vh, vals[k], font=r.hf, size=vs, color=_ink(t, k))
                yy = top + pad + vh + 0.02
                lab = plain(it.get("label") or (it.get("text") if it.get("head") else ""))
                if lab:
                    lh = _h(r, [(lab, True)], size, cw_r - 2 * pad)
                    text(s, cx + pad, yy, cw_r - 2 * pad, lh + 0.03, lab, font=r.bf_bold, size=size, color=t["text"])
                    yy += lh
                if it.get("desc"):
                    text(s, cx + pad, yy, cw_r - 2 * pad, cy + rh2 - yy, _to_paras(rich(it["desc"]), r.bf, r.bf_bold,
                                                                                   t["muted"]),
                         font=r.bf, size=size - 1, color=t["muted"], ls=LS)
    return total


def _body_fields(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    kw = min(w * 0.36, max(text_w(plain(i.get("head") or ""), r.bf_bold, size, True) for i in sec["items"]) + 0.25)
    gap = size / 72 * 0.55
    hs = [max(_h(r, rich(i.get("text") or ""), size, w - kw), _h(r, [(plain(i.get("head") or ""), True)], size, kw - 0.15))
          for i in sec["items"]]
    total = sum(hs) + gap * (len(hs) - 1)
    if draw:
        gap += min(0.25, extra / max(1, len(hs) - 1)) if len(hs) > 1 else 0
        yy = y
        for i, hh in zip(sec["items"], hs):
            text(s, x, yy, kw - 0.15, hh + 0.03, plain(i.get("head") or ""), font=r.bf_bold, size=size, color=t["accent_text"])
            text(s, x + kw, yy, w - kw, hh + 0.03, _to_paras(rich(i.get("text") or ""), r.bf, r.bf_bold, t["text"]),
                 font=r.bf, size=size, color=t["text"], ls=LS)
            yy += hh + gap
    return total


def _body_paragraph(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    runs = []
    for i in sec["items"]:
        runs += _item_runs(i) + [(" ", False)]
    runs = runs[:-1] if runs else [("", False)]
    h = _h(r, runs, size, w, 1.1)
    if draw:
        text(s, x, y, w, h + 0.05, _to_paras(runs, r.bf, r.bf_bold, t["text"]), font=r.bf, size=size, color=t["text"], ls=1.1)
    return h


def _body_table(r, s, sec, x, y, w, size, draw, extra=0.0):
    t = r.t
    header, rows = sec["header"], sec["rows"]
    ncol = max([len(header)] + [len(rw) for rw in rows]) if (header or rows) else 0
    if not ncol:
        return 0
    header = (header + [""] * ncol)[:ncol]
    rows = [(rw + [""] * ncol)[:ncol] for rw in rows]
    lens = [min(max(max([len(plain(header[c]))] + [len(plain(rw[c])) for rw in rows]), 6), 50) for c in range(ncol)]
    cws = [w * L / sum(lens) for L in lens]
    has_head = any(h.strip() for h in header)
    rh = [max(_h(r, rich(rw[c]), size, cws[c] - 0.16) for c in range(ncol)) + 0.12 for rw in rows]
    hh = (max(_h(r, [(plain(h), True)], size, cws[c] - 0.16) for c, h in enumerate(header)) + 0.14) if has_head else 0
    total = sum(rh) + hh
    if draw:
        rh = [h + min(0.35, extra / max(1, len(rh))) for h in rh]      # taller rows fill the zone
        n_rows = len(rows) + (1 if has_head else 0)
        shp = s.shapes.add_table(n_rows, ncol, E(x), E(y), E(w), E(total))
        tbl = shp.table
        tbl._tbl.tblPr.set("bandRow", "0")
        for c in range(ncol):
            tbl.columns[c].width = E(cws[c])
        all_rows = ([header] if has_head else []) + rows
        heights = ([hh] if has_head else []) + rh
        for ri, (vals, hgt) in enumerate(zip(all_rows, heights)):
            tbl.rows[ri].height = E(hgt)
            is_head = has_head and ri == 0
            for c in range(ncol):
                cell = tbl.cell(ri, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(_pill_color(t) if is_head else (_tints(t)[0] if ri % 2 else t["surface"] if not t["dark"] else t["surface2"]))
                cell.margin_left = cell.margin_right = E(0.08)
                cell.margin_top = cell.margin_bottom = E(0.05)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = ""
                p = tf.paragraphs[0]
                for seg, b in ([(plain(vals[c]), True)] if is_head else rich(vals[c])):
                    run = p.add_run()
                    run.text = seg
                    run.font.size = Pt(size)
                    run.font.bold = b or is_head or c == 0
                    run.font.name = r.bf_bold if (b or is_head or c == 0) else r.bf
                    run.font.color.rgb = _rgb(on_color(_pill_color(t)) if is_head else t["text"])
    return total


_BODIES = {"cards": _body_cards, "list": _body_list, "steps": _body_steps, "stats": _body_stats,
           "fields": _body_fields, "paragraph": _body_paragraph, "table": _body_table}


def _section_h(r, s, sec, x, y, w, size, draw, k0=0, extra=0.0):  # noqa: C901
    """Pill + optional section lead + body."""
    t = r.t
    h = 0.0
    if sec["heading"]:
        if draw:
            _pill(r, s, x, y, sec["heading"])
        h += PILL_H + 0.12
    if sec["lead"]:
        lh = _h(r, rich(sec["lead"]), size, w)
        if draw:
            text(s, x, y + h, w, lh + 0.04, _to_paras(rich(sec["lead"]), r.bf, r.bf_bold, t["muted"], t["text"]),
                 font=r.bf, size=size, color=t["muted"], ls=LS)
        h += lh + 0.08
    fn = _BODIES.get(sec["style"], _body_cards)
    if fn is _body_cards:
        h += fn(r, s, sec, x, y + h, w, size, draw, k0=k0, extra=extra)
    else:
        h += fn(r, s, sec, x, y + h, w, size, draw, extra=extra)
    return h


def _stretch_cap(sec) -> float:
    """How much extra height a section can absorb before it looks inflated."""
    n = max(1, len(sec["items"]) or len(sec["rows"]))
    return {"cards": 0.55, "stats": 0.6, "steps": 0.8, "table": 0.3, "list": 0.25, "fields": 0.22,
            "paragraph": 0.0}.get(sec["style"], 0.3) * (n if sec["style"] in ("list", "fields", "table") else
                                                          max(1, -(-n // max(1, _cols_for(n, 6)))))


# ══════════════════════════════════════════════════════════════════════════════
#  chrome
# ══════════════════════════════════════════════════════════════════════════════
def footer_band(r, s, idx):
    t = r.t
    c = _mix(t["accent"], "000000", 0.12) if not t["dark"] else t["surface2"]
    box(s, 0, SH - FOOT_H, SW, FOOT_H, fill=c)
    label = (r.deck.get("footer") or r.deck.get("title") or "")[:90]
    text(s, 1.5, SH - FOOT_H, SW - 3, FOOT_H, plain(label), font=r.bf, size=11, color=on_color(c), align="c", anchor="m")
    text(s, SW - MX - 0.8, SH - FOOT_H, 0.8, FOOT_H, str(idx), font=r.bf_bold, size=12, color=on_color(c),
         align="r", anchor="m")
    if r.logo:
        p, a = r._img(r.logo)
        if p:
            hh = 0.5
            place_image(s, p, SW - MX - min(hh * a, 2.2), 0.22, min(hh * a, 2.2), hh, a, mode="contain")


def _compact_header(r, s, sp, x, w, top=0.26):
    t = r.t
    title = plain(sp.get("title") or "")
    logo_room = 2.4 if r.logo else 0
    tw = w - logo_room
    box(s, x, top, 0.55, 0.06, fill=t["accent"])
    size = fit_size(title, r.hf, tw, 1.05, 28 if _COMPACT[0] or sp.get('_dense') else 34, 20, max_lines=2)
    th = para_h(title, r.hf, size, tw)
    text(s, x, top + 0.1, tw, th + 0.05, title, font=r.hf, size=size, color=t["text"])
    y = top + 0.1 + th
    sub = plain(sp.get("subtitle") or "")
    if sub:
        ss = fit_size(sub, r.bf, tw, 0.55, 16, 12, max_lines=2)
        sh_ = para_h(sub, r.bf, ss, tw)
        text(s, x, y + 0.04, tw, sh_ + 0.04, sub, font=r.bf, size=ss, color=t["muted"])
        y += 0.04 + sh_
    return y + 0.2


# ══════════════════════════════════════════════════════════════════════════════
#  image panel
# ══════════════════════════════════════════════════════════════════════════════
def _image_panel(r, s, imgs, captions, x, y, w, h, draw=True):
    """Framed images (no crop) + numbered captions under each. Returns used height."""
    t = r.t
    cap_size = 11
    aspects = [a for _, a in imgs]
    cap_h = 0.62 if captions else 0.0
    boxes = justified_rows(aspects, w, h, 0.2, cap_h=cap_h) if captions else justified_rows(aspects, w, h, 0.2)
    if not draw:
        return max(b[1] + b[3] for b in boxes) + cap_h
    _draw_image_boxes(r, s, imgs, captions, x, y, boxes)
    return max(b[1] + b[3] for b in boxes) + cap_h


def image_block(imgs, captions, w, h, rows):
    """Boxes for the images at full column width in exactly `rows` rows (scaled down only if taller than h).
    → (boxes relative to the block, block height incl. captions) or (None, 0)."""
    from app.services.ppt_designer import _justified_fixed
    cap_h = 0.62 if captions else 0.0
    aspects = [a for _, a in imgs]
    boxes = _justified_fixed(aspects, w, max(0.5, h - rows * cap_h), 0.2, rows)
    if not boxes:
        return None, 0.0
    ys = sorted({round(b[1], 4) for b in boxes})
    boxes = [(bx, by + ys.index(round(by, 4)) * cap_h, bw, bh) for bx, by, bw, bh in boxes]
    return boxes, max(b[1] + b[3] for b in boxes) + cap_h


def _draw_image_boxes(r, s, imgs, captions, x, y, boxes):
    t = r.t
    cap_size = 11
    for k, ((p, a), (bx, by, bw, bh)) in enumerate(zip(imgs, boxes)):
        pic = place_image(s, p, x + bx, y + by, bw, bh, a, mode="contain", radius=0.08, shadow=not t["dark"])
        if pic is not None:
            pic.line.color.rgb = _rgb(t["line"] if not t["dark"] else _mix(t["line"], "FFFFFF", 0.2))
            pic.line.width = Pt(1)
        if k < len(captions):
            cap = captions[k]
            cy = y + by + bh + 0.08
            d = 0.26
            c = _ink(t, k)
            box(s, x + bx, cy + 0.02, d, d, fill=c, shape=MSO_SHAPE.OVAL)
            text(s, x + bx, cy + 0.02, d, d, str(k + 1), font=r.bf_bold, size=10, color=on_color(c), align="c", anchor="m")
            paras = []
            if cap.get("head"):
                paras.append([(plain(cap["head"]), {"bold": True, "font": r.bf_bold, "color": t["text"], "size": cap_size + 1})])
            if cap.get("text"):
                paras.append([(plain(cap["text"]), {"color": t["muted"], "size": cap_size})])
            if paras:
                text(s, x + bx + d + 0.08, cy, max(0.8, bw - d - 0.08), 0.56, paras, font=r.bf, size=cap_size,
                     color=t["text"], ls=1.0)


# ══════════════════════════════════════════════════════════════════════════════
#  main composite slide
# ══════════════════════════════════════════════════════════════════════════════
def _masonry(heights: list[float], cols: int = 2) -> tuple[list[list[int]], float]:
    """Balanced split of sections into columns (reading order kept inside each column).
    Exhaustive for ≤ 8 sections, greedy beyond → (columns of indexes, tallest column)."""
    import itertools
    n = len(heights)

    def col_h(idx):
        return sum(heights[i] for i in idx) + SEC_GAP * max(0, len(idx) - 1)
    if n <= 8:
        best, best_key = None, None
        for assign in itertools.product(range(cols), repeat=n):
            if assign[0] != 0 or len(set(assign)) < min(cols, n):
                continue
            out = [[i for i in range(n) if assign[i] == c] for c in range(cols)]
            hs = [col_h(o) for o in out]
            key = (round(max(hs), 3), round(max(hs) - min(hs), 3))
            if best_key is None or key < best_key:
                best, best_key = out, key
        return best, best_key[0]
    colh = [0.0] * cols
    out = [[] for _ in range(cols)]
    for i, h in enumerate(heights):
        c = min(range(cols), key=lambda k: colh[k])
        out[c].append(i)
        colh[c] += h + (SEC_GAP if len(out[c]) > 1 else 0)
    return out, max(colh)


def _fill(sections: list, hs: list[float], room: float) -> tuple[list[float], float]:
    """Share the free height of a column between its sections (each up to its stretch cap).
    → (extra per section, extra gap between sections)."""
    free = room - sum(hs) - SEC_GAP * max(0, len(hs) - 1)
    if free <= 0 or not sections:
        return [0.0] * len(hs), 0.0
    caps = [_stretch_cap(x) for x in sections]
    extras = [0.0] * len(hs)
    left = free
    for _ in range(3):                                    # water-filling up to the caps
        open_ = [i for i in range(len(hs)) if extras[i] < caps[i] - 1e-6]
        if not open_ or left <= 1e-6:
            break
        share = left / len(open_)
        for i in open_:
            add = min(share, caps[i] - extras[i])
            extras[i] += add
            left -= add
    gap = min(0.3, left / max(1, len(hs) - 1)) if len(hs) > 1 else 0.0
    return extras, gap


def render_sections(r, s, sp, idx, total):
    t = r.t
    secs = _norm_sections(sp)
    imgs = r._images(sp)
    dense = r.deck.get("density") == "dense"
    _COMPACT[0] = False
    kinds = [x.get("kind") for x in r.deck.get("slides", [])]
    tight = kinds.count("sections") >= max(2, len(kinds) * 0.5)      # composite-heavy deck → reference grid
    CMX = globals()["CMX"] if tight else MX
    top = 0.26 if tight else MT
    if tight:
        footer_band(r, s, idx)
    else:
        r._footer(s, idx, total)
    cw = SW - 2 * CMX
    words = sum(len(plain(" ".join(i.get(k, "") for k in ("head", "text", "value", "label", "desc"))).split())
                for x in secs for i in x["items"]) + len(plain(sp.get("lead") or "").split())
    sp = {**sp, "_dense": words > 200}
    y0 = _compact_header(r, s, sp, CMX, cw, top)
    bottom = SH - FOOT_H - 0.1 if tight else SH - 0.62

    gallery = next((x for x in secs if x["style"] == "gallery"), None)
    want = (sp.get("image_section") or "").strip().lower()
    if want and imgs:                                  # the user said which section the images belong to
        named = next((x for x in secs if x["heading"] and (want in x["heading"].lower()
                                                           or x["heading"].lower() in want)), None)
        if named is not None:
            gallery = named
        elif gallery is None:
            gallery = {"heading": want.title(), "style": "gallery", "items": [], "highlight": False, "lead": "",
                       "header": [], "rows": []}
    if gallery is None and imgs:
        gallery = next((x for x in secs if _GALLERY.search(x["heading"]) and len(x["items"]) <= max(len(imgs), 1) + 1), None)
    # the bottom band is a row of cards: only a card-style group with 2–6 items can be it
    bandable = lambda x: x is not gallery and x["style"] in ("cards", "list", "paragraph") and 2 <= len(x["items"]) <= 6
    band = next((x for x in secs if x["highlight"] and bandable(x)), None)
    if band is None and len(secs) >= 3:
        cand = [x for x in secs if bandable(x) and _HIGHLIGHT.search(x["heading"])]
        band = cand[-1] if cand else None
    lead = sp.get("lead") or sp.get("body") or ""
    # de-duplicate: a one-item section that only restates the title/lead wastes a row;
    # an "Overview" one-liner *is* the lead (the user's own words), so promote it
    tok = lambda x: set(w[:6] for w in re.findall(r"[a-z0-9]{4,}", plain(x).lower()))
    first = (r.deck.get("slides") or [{}])[0]
    title_facts = " ".join([first.get("title") or "", first.get("subtitle") or ""] +
                           [str(i.get("text") or "") + " " + str(i.get("head") or "")
                            for x in (first.get("sections") or []) if isinstance(x, dict)
                            for i in (x.get("items") or []) if isinstance(i, dict)])
    known = tok(r.deck.get("title", "") + " " + (sp.get("title") or "") + " " + lead + " " + title_facts)
    kept = []
    for x in secs:
        if len(x["items"]) == 1 and x is not gallery and x is not band and x["style"] in ("cards", "fields", "list", "paragraph"):
            it = x["items"][0]
            words = tok(it.get("text") or "") or tok(it.get("head") or "")
            if re.search(r"overview|summary|about|introduction|description|what is", x["heading"], re.I) and \
                    len((it.get("text") or "").split()) >= 8:
                lead = it.get("text") or lead
                continue
            if words and len(words & known) / len(words) >= 0.7:
                continue
        kept.append(x)
    secs = kept
    base_stack = [x for x in secs if x is not gallery and x is not band]
    captions = gallery["items"][:len(imgs)] if gallery and imgs else []
    if gallery and not imgs:
        base_stack.append({**gallery, "style": "cards"})
    elif gallery and len(gallery["items"]) > len(imgs):
        # more screens described than images given → keep the rest as text right under the screenshots
        rest = gallery["items"][len(imgs):]
        base_stack.append({"heading": "", "style": "list" if all(not i.get("head") for i in rest) else "cards",
                           "items": rest, "highlight": False, "lead": "", "header": [], "rows": [],
                           "_with_gallery": True})
    max_s, min_s = (16, 10) if dense else (20, 11)

    def band_h(size):
        if not band:
            return 0.0
        n = len(band["items"])
        return PILL_H + 0.12 + _body_cards(r, s, band, CMX, 0, cw, size, False, cols=min(n, 4) if n != 5 else 5,
                                           equal=True)

    def evaluate(frac, mode, use_band, size, compact=False, img_min=1.4):
        """→ (fits, overflow, layout) for one candidate arrangement at one font size."""
        _COMPACT[0] = compact
        img_w = cw * frac if imgs else 0.0
        left_w = cw - (img_w + GAP if imgs else 0)
        stack = base_stack + ([] if use_band or not band else [band])
        bh = band_h(size) if (use_band and band) else 0.0
        lead_h = _lead_h(r, lead, size + 3, left_w) + 0.2 if lead else 0.0
        avail = bottom - y0 - (bh + 0.24 if bh else 0)
        if imgs and frac > 0:
            return eval_images(frac, use_band, size, compact, img_min, stack, left_w, img_w, bh, lead_h, avail)
        if mode.startswith("masonry") and len(stack) >= 2:
            ncol = 3 if mode == "masonry3" and len(stack) >= 3 else 2
            colw = (left_w - GAP * (ncol - 1)) / ncol
            hs = [_section_h(r, s, x, CMX, 0, colw, size, False) for x in stack]
            cols, stack_h = _masonry(hs, ncol)
        else:
            hs = [_section_h(r, s, x, CMX, 0, left_w, size, False) for x in stack]
            cols, stack_h = None, sum(hs) + SEC_GAP * max(0, len(hs) - 1)
        need = lead_h + stack_h
        fill = 1.0
        if imgs:                                   # images must stay a reasonable size and fill their panel
            ih = avail - (PILL_H + 0.14 if gallery and gallery["heading"] else 0)
            used = _image_panel(r, s, imgs, captions, 0, 0, img_w, ih, draw=False) if ih > 0.8 else 0
            if ih < img_min or used < img_min * 0.7:
                return False, 99, None
            fill = min(1.0, used / max(ih, 0.01))
        return need <= avail, need - avail, dict(frac=frac, img_w=img_w, left_w=left_w, stack=stack, cols=cols,
                                                 hs=hs, bh=bh, lead_h=lead_h, avail=avail, need=need, mode=mode,
                                                 use_band=use_band and bool(band), size=size, compact=compact,
                                                 fill=fill)

    def eval_images(frac, use_band, size, compact, img_min, stack, left_w, img_w, bh, lead_h, avail):
        """Right column = pill + images at full column width (natural aspect) + text sections under them.
        Sections are split between the columns so both end at the same height → no empty panel."""
        import itertools
        gp = PILL_H + 0.14 if gallery and gallery["heading"] else 0.0
        hsL = [_section_h(r, s, x, CMX, 0, left_w, size, False) for x in stack]
        hsR = [_section_h(r, s, x, CMX, 0, img_w, size, False) for x in stack]
        col = lambda hs, idx: sum(hs[i] for i in idx) + SEC_GAP * max(0, len(idx) - 1)
        best = None
        for rows in range(1, len(imgs) + 1):
            boxes, blk = image_block(imgs, captions, img_w, avail - gp, rows)
            if not boxes or min(b[3] for b in boxes) < img_min * 0.55:
                continue
            blk += gp
            for assign in itertools.product((0, 1), repeat=len(stack)):
                if any(stack[i].get("_with_gallery") and assign[i] == 0 for i in range(len(stack))):
                    continue                            # the gallery's own leftover items stay under the images
                L = [i for i, a_ in enumerate(assign) if a_ == 0]
                R = sorted((i for i, a_ in enumerate(assign) if a_ == 1),
                           key=lambda i: (0 if stack[i].get("_with_gallery") else 1, i))
                colL = lead_h + col(hsL, L)
                colR = blk + (SEC_GAP + col(hsR, R) if R else 0.0)
                need = max(colL, colR)
                area = sum(b[2] * b[3] for b in boxes)
                key = (need > avail + 1e-6, max(0.0, need - avail), abs(colL - colR) - 0.15 * area)
                if best is None or key < best[0]:
                    best = (key, dict(rows=rows, boxes=boxes, blk=blk, L=L, R=R, colL=colL, colR=colR, need=need,
                                      area=area))
        if best is None:
            return False, 99, None
        b_ = best[1]
        fill = 1.0 - min(1.0, abs(b_["colL"] - b_["colR"]) / max(avail, 0.1))
        fill += min(0.6, b_["area"] / 12.0)                # bigger screenshots are better
        return b_["need"] <= avail, b_["need"] - avail, dict(
            frac=frac, img_w=img_w, left_w=left_w, stack=stack, cols=None, hs=hsL, hsR=hsR, bh=bh, lead_h=lead_h,
            avail=avail, need=b_["need"], mode="images", use_band=use_band and bool(band), size=size,
            compact=compact, fill=fill, img=b_)

    fracs = [0.0] if not imgs else ([0.62] if not base_stack and not lead else
                                    [0.30, 0.34, 0.38, 0.42, 0.46, 0.50, 0.54, 0.58])
    many = sum(len(x["items"]) for x in base_stack)
    modes = ["single"]
    if not imgs and (len(base_stack) >= 2 and many >= 6 or len(base_stack) >= 3):
        modes += ["masonry"] + (["masonry3"] if len(base_stack) >= 4 else [])
    band_opts = [True, False] if band else [True]

    def search(img_min):
        best, best_key, fallback = None, None, None
        for compact in (False, True):
            for frac in fracs:
                for mode in modes:
                    for ub in band_opts:
                        for sz in _sizes(max_s, min_s - 1.5):
                            ok, over, lay_ = evaluate(frac, mode, ub, sz, compact, img_min)
                            if os.environ.get("PPT_DEBUG") == "2":
                                print(f"   try frac={frac} {mode} band={ub} compact={compact} size={sz}: "
                                      + (f"ok={ok} need={lay_['need']:.2f} avail={lay_['avail']:.2f}" if lay_ else "img-reject"))
                            if lay_ is None:
                                break
                            if ok:
                                # bigger text first; then keep the band, roomy cards, fewer columns, bigger images
                                # bigger text; images that fill their panel (no blank area under them)
                                key = (round(sz, 1) + (0.6 if ub else 0) - (0.8 if compact else 0)
                                       + {"single": 0.3, "masonry": 0.0, "masonry3": -0.4}.get(lay_["mode"], 0.0)
                                       + 2.5 * lay_["fill"], frac)
                                if best_key is None or key > best_key:
                                    best, best_key = lay_, key
                                break
                            if fallback is None or over < fallback[0]:
                                fallback = (over, lay_)
            if best and best["size"] >= min_s:
                break
        return best, fallback

    best, fallback = search(1.4)
    if best is None and imgs:
        best, fallback2 = search(0.9)
        fallback = fallback or fallback2
    lay = best or (fallback[1] if fallback else None)
    if lay is None:
        lay = evaluate(fracs[0], "single", True, min_s - 1.5, True, 0.5)[2] or evaluate(0.0, "single", True, min_s)[2]
    _COMPACT[0] = lay["compact"]
    size = lay["size"]
    if os.environ.get("PPT_DEBUG"):
        print(f"[composer] slide {idx}: fit={best is not None} size={size} frac={lay['frac']} mode={lay['mode']} "
              f"band={lay['use_band']} compact={lay['compact']} need={lay['need']:.2f} avail={lay['avail']:.2f} "
              f"lead={lay['lead_h']:.2f} hs={[round(h, 2) for h in lay['hs']]}")
    spare = max(0.0, lay["avail"] - lay["need"])

    # lead statement (bold, key words in accent) — spans the left zone
    y = y0
    left_w = lay["left_w"]
    if lead:
        paras = [[(seg, {"bold": True, "font": r.hf, "color": t["accent_text"] if b else t["text"]}) for seg, b in rich(lead)]]
        text(s, CMX, y, left_w, lay["lead_h"] - 0.15, paras, font=r.hf, size=size + 3, color=t["text"], ls=1.1)
        y += lay["lead_h"]
    stack = lay["stack"]
    region_bottom = bottom - (lay["bh"] + 0.24 if lay["bh"] else 0)
    if lay["mode"] == "images":
        im = lay["img"]
        k0 = 0
        secsL = [stack[i] for i in im["L"]]
        extras, gap_x = _fill(secsL, [lay["hs"][i] for i in im["L"]], region_bottom - y)
        yy = y
        for sec, ex in zip(secsL, extras):
            h = _section_h(r, s, sec, CMX, yy, left_w, size, True, k0=k0, extra=ex)
            k0 += len(sec["items"])
            yy += h + ex + SEC_GAP + gap_x
        ix, iy = CMX + left_w + GAP, y0
        secsR = [stack[i] for i in im["R"]]
        if not secsR:                                   # images alone: sit them optically centred in the column
            iy += max(0.0, (region_bottom - y0 - im["blk"]) * 0.35)
        if gallery and gallery["heading"]:
            _pill(r, s, ix, iy, gallery["heading"])
            iy += PILL_H + 0.14
        _draw_image_boxes(r, s, imgs, captions, ix, iy, im["boxes"])
        iy = y0 + im["blk"] + SEC_GAP if secsR else iy
        extras, gap_x = _fill(secsR, [lay["hsR"][i] for i in im["R"]], region_bottom - iy)
        for sec, ex in zip(secsR, extras):
            h = _section_h(r, s, sec, ix, iy, lay["img_w"], size, True, k0=k0, extra=ex)
            k0 += len(sec["items"])
            iy += h + ex + SEC_GAP + gap_x
    cols = [] if lay["mode"] == "images" else (lay["cols"] or [list(range(len(stack)))])
    colw = (left_w - GAP * (len(cols) - 1)) / max(1, len(cols))
    k0 = 0
    for c, col in enumerate(cols):                # every column is stretched down to the same baseline
        secs_c = [stack[i] for i in col]
        hs_c = [lay["hs"][i] for i in col]
        extras, gap_x = _fill(secs_c, hs_c, region_bottom - y)
        yy = y
        for sec, ex in zip(secs_c, extras):
            h = _section_h(r, s, sec, CMX + c * (colw + GAP), yy, colw, size, True, k0=k0, extra=ex)
            k0 += len(sec["items"])
            yy += h + ex + SEC_GAP + gap_x

    # image panel (right) — legacy path when the planner could not place images as a block
    if imgs and lay["img_w"] > 0 and lay["mode"] != "images":
        ix = CMX + left_w + GAP
        iy = y0
        ih = (bottom - (lay["bh"] + 0.24 if lay["bh"] else 0)) - iy
        if gallery and gallery["heading"]:
            _pill(r, s, ix, iy, gallery["heading"])
            iy += PILL_H + 0.14
            ih -= PILL_H + 0.14
        _image_panel(r, s, imgs, captions, ix, iy, lay["img_w"], ih)

    # bottom highlight band
    if lay["use_band"]:
        by = bottom - lay["bh"]
        pw = _pill(r, s, CMX, by, band["heading"] or "Highlights", color=t["text"] if not t["dark"] else t["accent"])
        if band["lead"]:
            lw = cw - pw - 0.3
            ls_ = fit_size(plain(band["lead"]), r.hf, lw, PILL_H + 0.1, size + 2, 10, max_lines=2)
            paras = [[(seg, {"bold": True, "font": r.hf, "color": t["accent_text"] if b else t["text"]})
                      for seg, b in rich(band["lead"])]]
            text(s, CMX + pw + 0.3, by - 0.05, lw, PILL_H + 0.15, paras, font=r.hf, size=ls_, color=t["text"], anchor="m")
        n = len(band["items"])
        _body_cards(r, s, band, CMX, by + PILL_H + 0.12, cw, size, True, k0=3, cols=min(n, 4) if n != 5 else 5, equal=True)


# ══════════════════════════════════════════════════════════════════════════════
#  title slide with details (event / problem statement / team fields)
# ══════════════════════════════════════════════════════════════════════════════
def render_title_sections(r, s, sp, idx, total):
    t = r.t
    secs = _norm_sections(sp)
    imgs = r._images(sp)
    c = _mix(t["accent"], "000000", 0.12) if not t["dark"] else t["surface2"]
    box(s, 0, SH - FOOT_H, SW, FOOT_H, fill=c)
    # decorative right panel
    panel_x = SW * 0.5
    box(s, panel_x, 0, SW - panel_x, SH - FOOT_H, fill=t["surface2"] if not t["dark"] else t["surface"])
    box(s, panel_x, 0, 0.1, SH - FOOT_H, fill=t["accent"])
    lx, lw = MX + 0.1, panel_x - MX - 0.6
    if r.logo:
        p, a = r._img(r.logo)
        if p:
            place_image(s, p, lx, 0.5, min(0.6 * a, 2.6), 0.6, a, mode="contain")
    title = plain(sp.get("title") or r.deck.get("title", ""))
    sub = plain(sp.get("subtitle") or "")
    lead = sp.get("lead") or sp.get("body") or ""
    ts = fit_size(title, r.hf, lw, 2.4, 54, 28, max_lines=3)
    th = para_h(title, r.hf, ts, lw)
    ss = fit_size(sub, r.bf, lw, 1.0, 22, 14, max_lines=3) if sub else 0
    sh_ = para_h(sub, r.bf, ss, lw) if sub else 0
    lsz = 14
    lh_ = _h(r, rich(lead), lsz, lw, 1.1) if lead else 0
    block = 0.3 + th + (0.2 + sh_ if sub else 0) + (0.3 + lh_ if lead else 0)
    y = max(1.3, (SH - FOOT_H - block) / 2)
    box(s, lx, y, 0.8, 0.09, fill=t["accent"])
    y += 0.3
    text(s, lx, y, lw, th + 0.05, title, font=r.hf, size=ts, color=t["text"])
    y += th + 0.2
    if sub:
        text(s, lx, y, lw, sh_ + 0.05, sub, font=r.bf, size=ss, color=t["accent_text"])
        y += sh_ + 0.3
    if lead:
        text(s, lx, y, lw, lh_ + 0.05, _to_paras(rich(lead), r.bf, r.bf_bold, t["muted"], t["text"]),
             font=r.bf, size=lsz, color=t["muted"], ls=1.1)
    # right: details card(s) or image
    rx, rw = panel_x + 0.55, SW - panel_x - 0.55 - MX
    top, bot = 0.7, SH - FOOT_H - 0.45
    if imgs and not secs:
        p, a = imgs[0]
        place_image(s, p, rx, top, rw, bot - top, a, mode="contain", radius=0.1, shadow=True)
        return
    _COMPACT[0] = False
    secs = [{**x, "style": "fields" if all(i.get("head") for i in x["items"]) else "list"} if x["style"] == "cards" else x
            for x in secs]
    size = 16
    for sz in _sizes(18, 10):
        need = sum(_section_h(r, s, x, rx + 0.3, 0, rw - 0.6, sz, False) for x in secs) + 0.25 * (len(secs) - 1) + 0.6
        if need <= bot - top:
            size = sz
            break
    need = sum(_section_h(r, s, x, rx + 0.3, 0, rw - 0.6, size, False) for x in secs) + 0.25 * (len(secs) - 1) + 0.6
    cy = top + max(0.0, (bot - top - need) / 2)
    box(s, rx, cy, rw, need, fill=t["bg"] if not t["dark"] else t["surface2"], radius=0.14, shadow=not t["dark"],
        line=t["line"])
    yy = cy + 0.3
    for x in secs:
        if x["style"] == "cards":
            x = {**x, "style": "fields" if all(i.get("head") for i in x["items"]) else "list"}
        yy += _section_h(r, s, x, rx + 0.3, yy, rw - 0.6, size, True) + 0.25
